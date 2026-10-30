"""Mobil API: giriş, rol yetkileri, bölüm kısıtı, onay akışı, fazla çalışma, iş talebi, çakışma ve işlem izi.
python -m unittest tests.test_api"""
import base64
import sys
import tempfile
import unittest
from datetime import date, timedelta
from io import BytesIO
from pathlib import Path
from urllib.parse import quote

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from fastapi.testclient import TestClient  # noqa: E402

from dose import db as dbm  # noqa: E402
from dose import denetim_formu  # noqa: E402
from dose import mevzuat as M  # noqa: E402
from dose import onay  # noqa: E402
from dose.db import Veritabani, sifre_hashle  # noqa: E402
from sunucu import api  # noqa: E402

YIL = 2026
FC = quote("Fazla Çalışma")


class MobilApi(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.yol = Path(cls.tmp.name) / "api.db"
        cls.yedekler = Path(cls.tmp.name) / "yedekler"
        cls.yedekler.mkdir()
        cls.onceki_api_yedekler, cls.onceki_db_yedekler = api.YEDEKLER, dbm.YEDEKLER
        api.YEDEKLER = dbm.YEDEKLER = cls.yedekler
        cls.app = api.uygulama_olustur(cls.yol)
        cls.c = TestClient(cls.app)
        d = cls.db = dbm.al()
        d.guncelle("kurum", 1, {"okul_adi": "Derik Mesleki ve Teknik Anadolu Lisesi",
                                "harcama_yetkilisi": "Ali Müdür", "kurum_kodu": "765432"})
        d.parametre_yaz(YIL, "fazla_calisma_saat_ucreti", 67.5)
        cls.bolum = d.ekle("bolum", {"ad": "Gıda Teknolojisi"})
        cls.bolum2 = d.ekle("bolum", {"ad": "Bilişim"})
        P = lambda ad, gorev, b: d.ekle("personel", {"ad_soyad": ad, "dose_gorevi": gorev, "bolum_id": b,
                                                      "hizmet_sinifi": "Eğitim ve Öğretim Hizmetleri"})
        cls.p_ogr = P("Deniz Öğretmen", "Öğretmen", cls.bolum)
        cls.p_ogr2 = P("Ece Öğretmen", "Öğretmen", cls.bolum)
        cls.p_sef = P("Can Şef", "Bölüm/Alan Şefi", cls.bolum)
        cls.p_bilisim = P("Bora Bilişim", "Öğretmen", cls.bolum2)
        for ad, rol, pid in (("deniz", M.ROL_OGRETMEN, cls.p_ogr), ("can", M.ROL_SEF, cls.p_sef),
                             ("veli", M.ROL_TMY, None), ("ali", M.ROL_MUDUR, None), ("ayse", M.ROL_SAYMAN, None),
                             ("deneme", M.ROL_IZLEYICI, None)):
            h, tuz = sifre_hashle("Gizli123")
            d.ekle("kullanici", {"kullanici_adi": ad, "ad_soyad": ad.title(), "rol": rol, "sifre_hash": h, "tuz": tuz,
                                 "personel_id": pid})
        h, tuz = sifre_hashle("Gizli123")
        cls.sistem_id = d.ekle("kullanici", {"kullanici_adi": "sistem", "ad_soyad": "Sistem Yöneticisi",
                                             "rol": M.ROL_SISTEM, "sifre_hash": h, "tuz": tuz,
                                             "sifre_degismeli": 0})
        d.guncelle("kullanici", d.deger("SELECT id FROM kullanici WHERE kullanici_adi = 'deniz'"), {"ad_soyad": "Deniz Öğretmen"})
        cls.musteri = d.ekle("cari", {"unvan": "Derik Belediyesi", "tur": "Müşteri", "kamu_kurumu": 1})
        cls.urun = d.ekle("urun", {"ad": "Zeytinyağı 1 L", "birim": "Adet", "satis_fiyati": 400, "bolum_id": cls.bolum})
        d.ekle("maliyet_kalem", {"urun_id": cls.urun, "kalem_turu": "Direkt İlk Madde ve Malzeme", "miktar": 5, "birim_fiyat": 40})
        d.ekle("maliyet_kalem", {"urun_id": cls.urun, "kalem_turu": "Ambalaj / Etiket", "miktar": 1, "birim_fiyat": 20})
        onay.akis_ayarla(True)

    @classmethod
    def tearDownClass(cls):
        cls.c.close()
        dbm.al().kapat()
        api.YEDEKLER = cls.onceki_api_yedekler
        dbm.YEDEKLER = cls.onceki_db_yedekler
        cls.tmp.cleanup()

    def giris(self, ad, sifre="Gizli123", cihaz="Test telefonu", kurum_kodu="765432") -> dict:
        govde = {"kullanici_adi": ad, "sifre": sifre, "cihaz": cihaz}
        if kurum_kodu:
            govde["kurum_kodu"] = kurum_kodu
        r = self.c.post("/api/giris", json=govde)
        self.assertEqual(r.status_code, 200, r.text)
        return {"Authorization": "Bearer " + r.json()["token"]}

    def test_parca_basi_api_hesaplama_kayit_belge_odeme_ve_bolum_yetkilerini_sunar(self):
        yetkili = self.giris("ayse")
        siparis_id = self.db.ekle("siparis", {
            "siparis_no": "PB-API-TEST", "tarih": f"{YIL}-03-01", "bolum_id": self.bolum,
            "konu": "API parça başı testi", "ucret_esasi": "Parça Başı Ücret (Yön. md.8)",
            "sozlesme_bedeli": 1000, "kdv_orani": 0.2, "iscilik_tutari": 700,
            "kabul_tarihi": f"{YIL}-03-15", "durum": "Tamamlandı (Geçici Kabul)",
        })
        body = {
            "yil": YIL, "ay": 3, "siparis_id": siparis_id, "kabul_edilen_iscilik": 100,
            "ogrenci_orani": 0.3, "aciklama": "API testi",
            "satirlar": [{"kisi_turu": "Personel", "kisi_id": self.p_ogr, "katki": 1,
                          "onceki_kumulatif": None, "puantaj": {"15": 1}}],
        }
        try:
            secenekler = self.c.get(f"/api/parca-basi/secenekler?yil={YIL}", headers=yetkili)
            self.assertEqual(secenekler.status_code, 200, secenekler.text)
            self.assertIn(siparis_id, [s["id"] for s in secenekler.json()["siparisler"]])
            self.assertIn(self.p_ogr, [p["id"] for p in secenekler.json()["personeller"]])

            onizleme = self.c.post("/api/parca-basi/hesapla", headers=yetkili, json=body)
            self.assertEqual(onizleme.status_code, 200, onizleme.text)
            self.assertAlmostEqual(onizleme.json()["hesap"]["toplamlar"]["hesaplanan"], 70)
            self.assertTrue(onizleme.json()["hesap"]["uyarilar"])

            ekle = self.c.post("/api/parca-basi", headers=yetkili, json=body)
            self.assertEqual(ekle.status_code, 200, ekle.text)
            pb_id = ekle.json()["id"]
            detay = self.c.get(f"/api/parca-basi/{pb_id}", headers=yetkili)
            self.assertEqual(detay.status_code, 200, detay.text)
            self.assertEqual(detay.json()["satirlar"][0]["kisi_id"], self.p_ogr)
            self.assertEqual(detay.json()["satirlar"][0]["puantaj"], {"15": 1})

            liste = self.c.get(f"/api/parca-basi?yil={YIL}&ay=3", headers=yetkili)
            self.assertEqual(liste.status_code, 200, liste.text)
            self.assertEqual(liste.json()["kayitlar"][0]["id"], pb_id)

            guncel = {**body, "kabul_edilen_iscilik": 120, "surum": self.db.son_islem_no()}
            duzenle = self.c.put(f"/api/parca-basi/{pb_id}", headers=yetkili, json=guncel)
            self.assertEqual(duzenle.status_code, 200, duzenle.text)
            self.assertEqual(self.db.deger("SELECT kabul_edilen_iscilik FROM parca_basi WHERE id = ?", (pb_id,)), 120)

            import base64
            from io import BytesIO
            import zipfile
            for tur in ("bordro", "puantaj", "banka"):
                belge = self.c.get(f"/api/parca-basi/{pb_id}/belge/{tur}", headers=yetkili)
                self.assertEqual(belge.status_code, 200, belge.text)
                icerik = base64.b64decode(belge.json()["icerik"])
                if tur == "banka":
                    self.assertTrue(zipfile.is_zipfile(BytesIO(icerik)))
                else:
                    self.assertTrue(icerik.startswith(b"%PDF"))

            self.db.guncelle("parca_basi", pb_id, {"durum": "Onaylandı"})
            odeme = self.c.post(f"/api/parca-basi/{pb_id}/odeme", headers=yetkili, json={
                "tarih": f"{YIL}-03-31", "hesap": "Banka (102)", "belge_no": "PB-ODEME",
            })
            self.assertEqual(odeme.status_code, 200, odeme.text)
            self.assertEqual(self.db.deger("SELECT kaynak FROM kasa_banka WHERE kaynak_id = ?", (pb_id,)), "ParcaBasi")
            tekrar = self.c.post(f"/api/parca-basi/{pb_id}/odeme", headers=yetkili, json={
                "tarih": f"{YIL}-03-31", "hesap": "Banka (102)",
            })
            self.assertEqual(tekrar.status_code, 400)
            sil = self.c.delete(f"/api/parca-basi/{pb_id}", headers=yetkili)
            self.assertEqual(sil.status_code, 400)
            ogretmen = self.giris("deniz")
            gorunum = self.c.get(f"/api/parca-basi?yil={YIL}&ay=3", headers=ogretmen)
            self.assertEqual(gorunum.status_code, 200, gorunum.text)
            self.assertEqual(gorunum.json()["kayitlar"][0]["net"], self.db.deger(
                "SELECT net FROM parca_basi_satir WHERE parca_basi_id = ?", (pb_id,)))
        finally:
            self.db.ex("DELETE FROM kasa_banka WHERE kaynak = 'ParcaBasi' AND kaynak_id IN "
                       "(SELECT id FROM parca_basi WHERE siparis_id = ?)", (siparis_id,))
            for pb in self.db.q("SELECT id FROM parca_basi WHERE siparis_id = ?", (siparis_id,)):
                self.db.ex("DELETE FROM parca_basi_satir WHERE parca_basi_id = ?", (pb["id"],))
                self.db.sil("parca_basi", pb["id"])
                onay.sil("Parça Başı", pb["id"])
            self.db.sil("siparis", siparis_id)

    def test_kurum_kodu_ve_kesif(self):
        d = self.c.get("/api/durum").json()
        self.assertEqual(d["kurum_kodu"], "765432")
        self.assertFalse(d["coklu"])
        r = self.c.post("/api/kesif", json={"kurum_kodu": "765432"})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertTrue(r.json()["uygun"])
        self.assertEqual(self.c.post("/api/kesif", json={"kurum_kodu": "111111"}).status_code, 404)
        r = self.c.post("/api/giris", json={"kullanici_adi": "deniz", "sifre": "Gizli123", "cihaz": "t",
                                           "kurum_kodu": "111111"})
        self.assertEqual(r.status_code, 403)
        r = self.c.post("/api/giris", json={"kullanici_adi": "deniz", "sifre": "Gizli123", "cihaz": "t",
                                           "kurum_kodu": "765432"})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(r.json()["kurum_kodu"], "765432")
        ben = self.c.get("/api/ben", headers={"Authorization": "Bearer " + r.json()["token"]}).json()
        self.assertEqual(ben["kurum_kodu"], "765432")

    def test_lisans_etkinlestirme_oturumsuz_anahtari_dogrular(self):
        from cryptography.hazmat.primitives import serialization
        from cryptography.hazmat.primitives.asymmetric.ed25519 import Ed25519PrivateKey
        from unittest import mock

        private_key = Ed25519PrivateKey.generate()
        public_key = api.lisans._b64(private_key.public_key().public_bytes(
            serialization.Encoding.Raw, serialization.PublicFormat.Raw))
        anahtar = api.lisans.imzala(
            private_key, "Derik Mesleki ve Teknik Anadolu Lisesi", "765432",
            date.today() + timedelta(days=365), "TEST-LISANS", "1 yıl")

        try:
            with mock.patch.object(api.lisans, "ACIK_ANAHTAR", public_key):
                gecersiz = self.c.post("/api/lisans/etkinlestir", json={
                    "kurum_kodu": "765432", "anahtar": "gecersiz lisans anahtari ve imza",
                })
                self.assertEqual(gecersiz.status_code, 400, gecersiz.text)

                yanlis_kod = self.c.post("/api/lisans/etkinlestir", json={
                    "kurum_kodu": "111111", "anahtar": anahtar,
                })
                self.assertEqual(yanlis_kod.status_code, 403, yanlis_kod.text)

                etkinlestirildi = self.c.post("/api/lisans/etkinlestir", json={
                    "kurum_kodu": "765432", "anahtar": anahtar,
                })
                self.assertEqual(etkinlestirildi.status_code, 200, etkinlestirildi.text)
                self.assertEqual(etkinlestirildi.json()["lisans"], "TEST-LISANS")
                self.assertFalse(self.db.salt_okunur)
                durum = self.c.get("/api/durum").json()
                self.assertIn("Lisanslı: Derik Mesleki ve Teknik Anadolu Lisesi", durum["lisans"])
                self.assertIn((date.today() + timedelta(days=365)).strftime("%d.%m.%Y"), durum["lisans"])
        finally:
            self.db.ayar_yaz("lisans", "")
            self.db.salt_okunur = False
            api._son_lisans[0] = 0.0

    def test_tanim_api_yetki_validasyon_ve_kullanimda_silme_korumasini_uygular(self):
        admin = self.giris("ali")
        ekle = self.c.post("/api/tanimlar/bolum", headers=admin,
                           json={"degerler": {"ad": "Flet API Test Bölümü",
                                              "utp_yontemi": M.UTP_BOLUM_YONTEMLERI[0]}})
        self.assertEqual(ekle.status_code, 200, ekle.text)
        bolum_id = ekle.json()["id"]

        liste = self.c.get("/api/tanimlar/bolum?ara=Flet%20API", headers=admin)
        self.assertEqual(liste.status_code, 200, liste.text)
        self.assertTrue(liste.json()["yazabilir"])
        self.assertEqual(liste.json()["kayitlar"][0]["id"], bolum_id)
        self.assertIn("aktif", {a["anahtar"] for a in liste.json()["alanlar"]})

        eksik = self.c.post("/api/tanimlar/bolum", headers=admin, json={"degerler": {"ad": ""}})
        self.assertEqual(eksik.status_code, 400)
        gecersiz_alan = self.c.post("/api/tanimlar/bolum", headers=admin,
                                    json={"degerler": {"ad": "Geçersiz", "admin": True}})
        self.assertEqual(gecersiz_alan.status_code, 400)

        personel = self.c.post("/api/tanimlar/personel", headers=admin,
                               json={"degerler": {"ad_soyad": "Flet Test Personeli",
                                                  "hizmet_sinifi": M.HIZMET_SINIFLARI[0],
                                                  "dose_gorevi": "Öğretmen", "bolum_id": bolum_id}})
        self.assertEqual(personel.status_code, 200, personel.text)
        kullanilan_bolum = self.c.delete(f"/api/tanimlar/bolum/{bolum_id}", headers=admin)
        self.assertEqual(kullanilan_bolum.status_code, 409)

        gecersiz_tckn = self.c.post("/api/tanimlar/personel", headers=admin,
                                    json={"degerler": {"ad_soyad": "Geçersiz Kimlik",
                                                       "tckn": "123",
                                                       "hizmet_sinifi": M.HIZMET_SINIFLARI[0],
                                                       "dose_gorevi": "Öğretmen"}})
        self.assertEqual(gecersiz_tckn.status_code, 400)

        ogretmen = self.giris("deniz")
        yetkisiz = self.c.get("/api/tanimlar/bolum", headers=ogretmen)
        self.assertEqual(yetkisiz.status_code, 403)

        sil = self.c.delete(f"/api/tanimlar/personel/{personel.json()['id']}", headers=admin)
        self.assertEqual(sil.status_code, 200, sil.text)
        sil = self.c.delete(f"/api/tanimlar/bolum/{bolum_id}", headers=admin)
        self.assertEqual(sil.status_code, 200, sil.text)

    def test_urun_maliyet_api_hesaplama_kayit_fiyat_guncelleme_ve_tutanak_sunar(self):
        yetkili = self.giris("veli")
        urun_id = self.db.ekle("urun", {
            "ad": "Maliyet API testi", "birim": "Adet", "bolum_id": self.bolum,
            "ucret_esasi": M.UCRET_ESASLARI[0], "satis_fiyati": 500, "kar_orani": 0.1, "kdv_orani": 0.2,
        })
        govde = {
            "oto": False, "hedef_fiyat": 500, "kar_orani": 0.1, "kdv_orani": 0.2,
            "piyasa": "Piyasa araştırması örneği",
            "kalemler": [
                {"kalem_turu": "Direkt İlk Madde ve Malzeme", "aciklama": "Hammadde",
                 "miktar": 1, "birim_fiyat": 100},
                {"kalem_turu": "Direkt İşçilik (Personel)", "miktar": 1, "birim_fiyat": 200},
                {"kalem_turu": "Öğrenci Hakkı", "miktar": 1, "birim_fiyat": 50},
                {"kalem_turu": "Enerji (Elektrik/Gaz/Su)", "miktar": 1, "birim_fiyat": 30},
            ],
        }
        try:
            detay = self.c.get(f"/api/urun/{urun_id}/maliyet", headers=yetkili)
            self.assertEqual(detay.status_code, 200, detay.text)
            self.assertTrue(detay.json()["yazabilir"])
            self.assertIn("Ambalaj / Etiket", detay.json()["kalem_turleri"])

            hesap = self.c.post(f"/api/urun/{urun_id}/maliyet/hesapla", headers=yetkili, json=govde)
            self.assertEqual(hesap.status_code, 200, hesap.text)
            self.assertEqual(hesap.json()["ozet"]["birim_maliyet"], 380)
            self.assertEqual(hesap.json()["ozet"]["satis_fiyati"], 418)
            self.assertTrue(hesap.json()["uyarilar"])

            otomatik = self.c.post(
                f"/api/urun/{urun_id}/maliyet/hesapla", headers=yetkili,
                json={**govde, "oto": True, "hedef_fiyat": 1000},
            )
            self.assertEqual(otomatik.status_code, 200, otomatik.text)
            self.assertEqual(otomatik.json()["ozet"]["satis_fiyati"], 1000)
            self.assertTrue(any(k["kalem_turu"] == "Direkt İşçilik (Personel)" and k["birim_fiyat"] > 0
                                for k in otomatik.json()["kalemler"]))

            hedef = {**govde, "hedef_dahil": 600}
            hedef_hesabi = self.c.post(f"/api/urun/{urun_id}/maliyet/hesapla", headers=yetkili, json=hedef)
            self.assertEqual(hedef_hesabi.status_code, 200, hedef_hesabi.text)
            self.assertAlmostEqual(hedef_hesabi.json()["kar_orani"], (600 / 1.2) / 380 - 1)

            kayit = self.c.put(f"/api/urun/{urun_id}/maliyet", headers=yetkili, json=govde)
            self.assertEqual(kayit.status_code, 200, kayit.text)
            self.assertEqual(self.db.deger("SELECT COUNT(*) FROM maliyet_kalem WHERE urun_id = ?", (urun_id,)), 4)

            fiyat = self.c.post(f"/api/urun/{urun_id}/maliyet/fiyat", headers=yetkili, json=govde)
            self.assertEqual(fiyat.status_code, 200, fiyat.text)
            self.assertEqual(self.db.deger("SELECT satis_fiyati FROM urun WHERE id = ?", (urun_id,)), 418)

            tutanak = self.c.post(f"/api/urun/{urun_id}/maliyet/tutanak", headers=yetkili, json=govde)
            self.assertEqual(tutanak.status_code, 200, tutanak.text)
            self.assertTrue(base64.b64decode(tutanak.json()["icerik"]).startswith(b"%PDF"))
            self.assertEqual(self.db.deger("SELECT COUNT(*) FROM fiyat_tespit WHERE urun_id = ?", (urun_id,)), 1)

            okuyucu = self.giris("deneme")
            yazma_yetkisiz = self.c.put(f"/api/urun/{urun_id}/maliyet", headers=okuyucu, json=govde)
            self.assertEqual(yazma_yetkisiz.status_code, 403)
        finally:
            self.db.ex("DELETE FROM maliyet_kalem WHERE urun_id = ?", (urun_id,))
            self.db.ex("DELETE FROM fiyat_tespit WHERE urun_id = ?", (urun_id,))
            self.db.sil("urun", urun_id)

    def test_mevzuat_parametre_api_yetkili_yilda_tam_kayit_yapar(self):
        admin = self.giris("ayse")
        sonuc = self.c.get(f"/api/parametreler?yil={YIL}", headers=admin)
        self.assertEqual(sonuc.status_code, 200, sonuc.text)
        self.assertTrue(sonuc.json()["yazabilir"])
        degerler = {p["anahtar"]: p["deger"] for p in sonuc.json()["parametreler"]}
        kayit = self.c.put(f"/api/parametreler?yil={YIL}", headers=admin,
                           json={"degerler": degerler})
        self.assertEqual(kayit.status_code, 200, kayit.text)
        eksik = self.c.put(f"/api/parametreler?yil={YIL}", headers=admin,
                           json={"degerler": {}})
        self.assertEqual(eksik.status_code, 400)

    def test_mevzuat_rehberi_api_sunari_icerigi_takvimi_parametreleri_ve_salt_okunurlugu(self):
        okuyucu = self.giris("deneme")
        yanit = self.c.get(f"/api/rehber?yil={YIL}", headers=okuyucu)
        self.assertEqual(yanit.status_code, 200, yanit.text)
        rehber = yanit.json()
        self.assertEqual(rehber["yil"], YIL)
        self.assertIn("DÖSE İşlemleri Mevzuat Rehberi", rehber["mevzuat_html"])
        self.assertTrue(any(kayit["terim"] == "Harcama yetkilisi" for kayit in rehber["sozluk"]))
        self.assertTrue(all({"hata", "sonuc", "koruma", "dayanak"} <= kayit.keys()
                            for kayit in rehber["riskler"]))
        self.assertTrue(any(rol["unvan"].startswith("Sayman") for rol in rehber["roller"]))
        self.assertTrue(rehber["yasal_takvim"])
        self.assertEqual(rehber["merkez_saymanlik"]["vkn"], "6210312482")
        self.assertTrue(any(kayit["etiket"] == "Aylık brüt asgari ücret (TL)"
                            and kayit["deger"] == "33.030,00 TL" for kayit in rehber["parametreler"]))
        for sure in rehber["yaklasan_sureler"]:
            gun = (date.fromisoformat(sure["tarih"]) - date.today()).days
            self.assertGreaterEqual(gun, 0)
            self.assertLessEqual(gun, 60)

        yetkisiz = self.c.get(f"/api/rehber?yil={YIL}")
        self.assertEqual(yetkisiz.status_code, 401)

    def test_kurum_bilgisi_api_yetki_ve_alan_dogrulamasi(self):
        admin = self.giris("ali")
        sonuc = self.c.get("/api/kurum", headers=admin)
        self.assertEqual(sonuc.status_code, 200, sonuc.text)
        self.assertTrue(sonuc.json()["yazabilir"])
        guncelle = self.c.put("/api/kurum", headers=admin,
                              json={"degerler": {"okul_adi": "Flet Kurum Testi", "tahsis_sermaye": 2500,
                                                 "buyuksehir": True}})
        self.assertEqual(guncelle.status_code, 200, guncelle.text)
        self.assertEqual(self.db.kurum()["okul_adi"], "Flet Kurum Testi")
        hatali = self.c.put("/api/kurum", headers=admin,
                            json={"degerler": {"okul_adi": "", "vkn": "123"}})
        self.assertEqual(hatali.status_code, 400)
        self.assertEqual(self.db.kurum()["okul_adi"], "Flet Kurum Testi")
        self.db.guncelle("kurum", 1, {"okul_adi": "Derik Mesleki ve Teknik Anadolu Lisesi"})

    def test_muhasebe_api_yevmiye_mizan_gelir_ve_fis_detayini_sunar(self):
        sayman = self.giris("ayse")
        fis_id = self.db.ekle("yevmiye", {"fis_no": "MIF-TEST", "tarih": f"{YIL}-01-15",
                                          "aciklama": "API muhasebe testi", "kaynak": "Manuel"})
        self.db.ekle("yevmiye_satir", {"yevmiye_id": fis_id, "hesap_kodu": "100", "aciklama": "Kasa",
                                       "borc": 100, "alacak": 0})
        self.db.ekle("yevmiye_satir", {"yevmiye_id": fis_id, "hesap_kodu": "600", "aciklama": "Gelir",
                                       "borc": 0, "alacak": 100})
        ozet = self.c.get(f"/api/muhasebe?yil={YIL}", headers=sayman)
        self.assertEqual(ozet.status_code, 200, ozet.text)
        self.assertTrue(ozet.json()["yazabilir"])
        self.assertEqual([f["id"] for f in ozet.json()["fisler"]], [fis_id])
        self.assertEqual(ozet.json()["mizan"][0]["borc"], 100)
        self.assertEqual(ozet.json()["gelir_tablosu"]["toplam_gelir"], 100)
        detay = self.c.get(f"/api/muhasebe/yevmiye/{fis_id}", headers=sayman)
        self.assertEqual(detay.status_code, 200, detay.text)
        self.assertEqual(len(detay.json()["satirlar"]), 2)
        hatali = self.c.get(f"/api/muhasebe?baslangic=2026-02-02&bitis=2026-01-01",
                            headers=sayman)
        self.assertEqual(hatali.status_code, 400)
        with self.db.islem():
            self.db.ex("DELETE FROM yevmiye_satir WHERE yevmiye_id = ?", (fis_id,))
            self.db.sil("yevmiye", fis_id)

    def test_muhasebe_api_manuel_fis_kaydeder_dengeler_ve_yetkiyi_korur(self):
        sayman = self.giris("ayse")
        govde = {"tarih": f"{YIL}-02-01", "aciklama": "API manuel fiş",
                 "satirlar": [
                     {"hesap_kodu": "100", "aciklama": "Kasa", "borc": 250, "alacak": 0},
                     {"hesap_kodu": "600", "aciklama": "Gelir", "borc": 0, "alacak": 250},
                 ]}
        kayit = self.c.post("/api/muhasebe/yevmiye", headers=sayman, json=govde)
        self.assertEqual(kayit.status_code, 200, kayit.text)
        fis_id = kayit.json()["id"]
        self.assertEqual(self.db.q1("SELECT kaynak FROM yevmiye WHERE id = ?", (fis_id,))["kaynak"], "Manuel")
        dengesiz = {**govde, "satirlar": [
            {"hesap_kodu": "100", "borc": 251, "alacak": 0},
            {"hesap_kodu": "600", "borc": 0, "alacak": 250},
        ]}
        self.assertEqual(self.c.post("/api/muhasebe/yevmiye", headers=sayman, json=dengesiz).status_code, 400)
        gecersiz_hesap = {**govde, "satirlar": [
            {"hesap_kodu": "HESAP-YOK", "borc": 250, "alacak": 0},
            {"hesap_kodu": "600", "borc": 0, "alacak": 250},
        ]}
        self.assertEqual(self.c.post("/api/muhasebe/yevmiye", headers=sayman, json=gecersiz_hesap).status_code, 400)
        mudur = self.giris("ali")
        self.assertEqual(self.c.post("/api/muhasebe/yevmiye", headers=mudur, json=govde).status_code, 403)
        self.assertEqual(self.c.post("/api/muhasebe/yeniden-uret", headers=mudur).status_code, 403)
        silindi = self.c.delete(f"/api/muhasebe/yevmiye/{fis_id}", headers=sayman)
        self.assertEqual(silindi.status_code, 200, silindi.text)
        self.assertIsNone(self.db.q1("SELECT id FROM yevmiye WHERE id = ?", (fis_id,)))

    def test_muhasebe_api_hesap_plani_yazma_ve_kullanimda_silme_korumasini_uygular(self):
        sayman = self.giris("ayse")
        ekle = self.c.post("/api/muhasebe/hesap-plani", headers=sayman,
                           json={"kod": "TEST-900", "ad": "Test hesabı"})
        self.assertEqual(ekle.status_code, 200, ekle.text)
        tekrar = self.c.post("/api/muhasebe/hesap-plani", headers=sayman,
                             json={"kod": "TEST-900", "ad": "Tekrar"})
        self.assertEqual(tekrar.status_code, 409)

        fis = self.c.post("/api/muhasebe/yevmiye", headers=sayman, json={
            "tarih": f"{YIL}-02-02", "aciklama": "Hesap silme koruma testi",
            "satirlar": [
                {"hesap_kodu": "TEST-900", "borc": 50, "alacak": 0},
                {"hesap_kodu": "TEST-900", "borc": 0, "alacak": 50},
            ],
        })
        self.assertEqual(fis.status_code, 200, fis.text)
        fis_id = fis.json()["id"]
        koruma = self.c.delete("/api/muhasebe/hesap-plani/TEST-900", headers=sayman)
        self.assertEqual(koruma.status_code, 409)
        self.assertEqual(self.c.delete(f"/api/muhasebe/yevmiye/{fis_id}", headers=sayman).status_code, 200)
        sil = self.c.delete("/api/muhasebe/hesap-plani/TEST-900", headers=sayman)
        self.assertEqual(sil.status_code, 200, sil.text)
        self.assertIsNone(self.db.q1("SELECT kod FROM hesap_plani WHERE kod = 'TEST-900'"))

    def test_rapor_api_tum_masaustu_raporlarini_ve_disari_aktarimi_sunar(self):
        sayman = self.giris("ayse")
        raporlar = (
            "siparis", "urun", "bolum", "aylik", "cari", "personel", "ogrenci",
            "temin", "hazine", "uc_aylik", "idari_masraf",
        )
        for rapor_id in raporlar:
            with self.subTest(rapor=rapor_id):
                yanit = self.c.get(f"/api/raporlar?rapor={rapor_id}&yil={YIL}", headers=sayman)
                self.assertEqual(yanit.status_code, 200, yanit.text)
                veri = yanit.json()
                self.assertEqual(veri["id"], rapor_id)
                self.assertTrue(veri["kolonlar"])
                self.assertEqual(veri["ozet"]["satir_sayisi"], len(veri["satirlar"]))
                self.assertEqual(len(veri["secenekler"]), len(raporlar))
        bilinmeyen = self.c.get("/api/raporlar?rapor=bilinmeyen", headers=sayman)
        self.assertEqual(bilinmeyen.status_code, 404)
        yetkisiz = self.c.get("/api/raporlar", headers=self.giris("deniz"))
        self.assertEqual(yetkisiz.status_code, 403)

        excel = self.c.get(f"/api/raporlar/siparis/disa-aktar?yil={YIL}&bicim=xlsx", headers=sayman)
        self.assertEqual(excel.status_code, 200, excel.text)
        import base64
        import zipfile
        from io import BytesIO
        excel_bytes = base64.b64decode(excel.json()["icerik"])
        self.assertTrue(zipfile.is_zipfile(BytesIO(excel_bytes)))
        pdf = self.c.get(f"/api/raporlar/uc_aylik/disa-aktar?yil={YIL}&bicim=pdf", headers=sayman)
        self.assertEqual(pdf.status_code, 200, pdf.text)
        self.assertTrue(base64.b64decode(pdf.json()["icerik"]).startswith(b"%PDF"))
        self.assertEqual(self.c.get(f"/api/raporlar/siparis/disa-aktar?bicim=csv",
                                    headers=sayman).status_code, 422)

    def test_vergi_api_kdv_devreden_muhtasar_ayrinti_ve_excel_sunar(self):
        sayman = self.giris("ayse")
        fatura_ids = []
        try:
            for tur, ay, matrah, kdv, stopaj in (
                ("Satış", "01", 1000, 200, 0),
                ("Alış", "01", 2000, 400, 10),
                ("Satış", "02", 500, 100, 0),
            ):
                fatura_ids.append(self.db.ekle("fatura", {
                    "tur": tur, "belge_turu": "Diğer", "belge_no": f"VERGI-{tur}-{ay}",
                    "tarih": f"{YIL}-{ay}-15", "matrah": matrah, "kdv": kdv,
                    "stopaj_turu": "Zirai Ürün", "stopaj": stopaj,
                    "tevkif_kdv": 0, "iptal": 0,
                }))
            yanit = self.c.get(f"/api/vergi?yil={YIL}", headers=sayman)
            self.assertEqual(yanit.status_code, 200, yanit.text)
            ocak, subat = yanit.json()["aylar"][:2]
            self.assertEqual((ocak["hesaplanan"], ocak["indirilecek"], ocak["odenecek_kdv"], ocak["devreden"]),
                             (200, 400, 0, 200))
            self.assertEqual((subat["hesaplanan"], subat["devreden"]), (100, 100))
            self.assertEqual(ocak["stopaj"], 10)
            detay = self.c.get(f"/api/vergi/{YIL}/1/muhtasar", headers=sayman)
            self.assertEqual(detay.status_code, 200, detay.text)
            self.assertEqual(detay.json()["detay"][0]["tutar"], 10)
            hatali_ay = self.c.get(f"/api/vergi/{YIL}/13/muhtasar", headers=sayman)
            self.assertEqual(hatali_ay.status_code, 400)
            excel = self.c.get(f"/api/vergi/disa-aktar?yil={YIL}&bicim=xlsx", headers=sayman)
            self.assertEqual(excel.status_code, 200, excel.text)
            import base64
            import zipfile
            from io import BytesIO
            self.assertTrue(zipfile.is_zipfile(BytesIO(base64.b64decode(excel.json()["icerik"]))))
        finally:
            with self.db.islem():
                for fatura_id in fatura_ids:
                    self.db.ex("DELETE FROM fatura_kalem WHERE fatura_id = ?", (fatura_id,))
                    self.db.sil("fatura", fatura_id)

    def test_hazine_yukumluluk_api_hesap_bildirim_odeme_ve_yetki_kontrolu(self):
        sayman = self.giris("ayse")
        kasa_id = self.db.ekle("kasa_banka", {
            "tarih": f"{YIL}-01-20", "hesap": "Banka (102)", "yon": "Giriş",
            "tur": "Tahsilat", "tutar": 500, "hasilat_tutari": 500,
            "kaynak": None, "kaynak_id": None,
        })
        hazine_id = None
        try:
            ozet = self.c.get(f"/api/yukumluluk/hazine?yil={YIL}", headers=sayman)
            self.assertEqual(ozet.status_code, 200, ozet.text)
            self.assertEqual(ozet.json()["aylik"][0]["guncel"], 500)
            hesapla = self.c.post("/api/yukumluluk/hazine/hesapla", headers=sayman,
                                  json={"yil": YIL, "ay": 1})
            self.assertEqual(hesapla.status_code, 200, hesapla.text)
            hazine = self.db.q1("SELECT * FROM hazine_payi WHERE yil = ? AND ay = ?", (YIL, 1))
            hazine_id = hazine["id"]
            self.assertEqual(hazine["tutar"], 5)
            pdf = self.c.get(f"/api/yukumluluk/hazine/{YIL}/1/pdf", headers=sayman)
            self.assertEqual(pdf.status_code, 200, pdf.text)
            import base64
            self.assertTrue(base64.b64decode(pdf.json()["icerik"]).startswith(b"%PDF"))
            excel = self.c.get(f"/api/yukumluluk/hazine/disa-aktar?yil={YIL}&bicim=xlsx", headers=sayman)
            self.assertEqual(excel.status_code, 200, excel.text)
            import zipfile
            from io import BytesIO
            self.assertTrue(zipfile.is_zipfile(BytesIO(base64.b64decode(excel.json()["icerik"]))))
            bildir = self.c.post("/api/yukumluluk/hazine/bildir", headers=sayman,
                                 json={"yil": YIL, "ay": 1})
            self.assertEqual(bildir.status_code, 200, bildir.text)
            odeme = self.c.post("/api/yukumluluk/odeme", headers=sayman,
                                json={"tur": "hazine", "yil": YIL, "ay": 1})
            self.assertEqual(odeme.status_code, 200, odeme.text)
            self.assertEqual(odeme.json()["tutar"], 5)
            tekrar = self.c.post("/api/yukumluluk/odeme", headers=sayman,
                                 json={"tur": "hazine", "yil": YIL, "ay": 1})
            self.assertEqual(tekrar.status_code, 409)
            state = self.db.q1("SELECT durum FROM hazine_payi WHERE id = ?", (hazine_id,))
            self.assertEqual(state["durum"], "Ödendi")
            ziyaretci = self.c.get(f"/api/yukumluluk/hazine?yil={YIL}", headers=self.giris("deniz"))
            self.assertEqual(ziyaretci.status_code, 403)
            salt_okunur = self.c.post("/api/yukumluluk/hazine/hesapla",
                                      headers=self.giris("deneme"), json={"yil": YIL, "ay": 2})
            self.assertEqual(salt_okunur.status_code, 403)
        finally:
            if hazine_id:
                with self.db.islem():
                    for fis in self.db.q("SELECT id FROM yevmiye WHERE kaynak = 'KasaBanka' AND kaynak_id IN "
                                         "(SELECT id FROM kasa_banka WHERE kaynak = 'HazinePayi' AND kaynak_id = ?)",
                                         (hazine_id,)):
                        self.db.ex("DELETE FROM yevmiye_satir WHERE yevmiye_id = ?", (fis["id"],))
                        self.db.sil("yevmiye", fis["id"])
                    for fis in self.db.q("SELECT id FROM yevmiye WHERE kaynak = 'HazinePayi' AND kaynak_id = ?",
                                         (hazine_id,)):
                        self.db.ex("DELETE FROM yevmiye_satir WHERE yevmiye_id = ?", (fis["id"],))
                        self.db.sil("yevmiye", fis["id"])
                    self.db.ex("DELETE FROM kasa_banka WHERE kaynak = 'HazinePayi' AND kaynak_id = ?", (hazine_id,))
                    self.db.sil("hazine_payi", hazine_id)
            self.db.sil("kasa_banka", kasa_id)

    def test_giris_kilit_ve_yetkiler(self):
        self.assertEqual(self.c.get("/api/durum").json()["satici"], "Derik MTAL Döner Sermaye İşletmesi")
        self.assertEqual(self.c.get("/api/panel").status_code, 401)
        r = self.c.post("/api/giris", json={"kullanici_adi": "ayse", "sifre": "yanlis", "kurum_kodu": "765432"})
        self.assertEqual((r.status_code, r.json()["hata"]), (401, "Kullanıcı adı veya şifre hatalı."))
        for _ in range(4):
            self.c.post("/api/giris", json={"kullanici_adi": "ayse", "sifre": "yanlis", "kurum_kodu": "765432"})
        r = self.c.post("/api/giris", json={"kullanici_adi": "ayse", "sifre": "Gizli123", "kurum_kodu": "765432"})
        self.assertEqual(r.status_code, 429)
        api._hatali_girisler.clear()

        ogr = self.giris("deniz")
        ben = self.c.get("/api/ben", headers=ogr).json()
        self.assertEqual(ben["moduller"]["fazla"], "y")
        self.assertEqual(ben["moduller"]["kasa"], "")
        self.assertEqual(ben["kullanici"]["bolum_id"], self.bolum)
        r = self.c.get("/api/kasa", headers=ogr)
        self.assertEqual(r.status_code, 403)
        self.assertIn("Kasa / Banka", r.json()["hata"])
        self.assertEqual(self.c.get("/api/islem-kaydi", headers=ogr).status_code, 403)
        izl = self.giris("deneme")
        self.assertEqual(self.c.get("/api/kasa", headers=izl).status_code, 200)
        r = self.c.post("/api/fazla-calisma", headers=izl, json={"yil": YIL, "ay": 3, "personel_id": self.p_ogr, "saat": 2})
        self.assertEqual(r.status_code, 403)
        self.c.post("/api/cikis", headers=izl)
        self.assertEqual(self.c.get("/api/ben", headers=izl).status_code, 401)

    def test_ilk_giriste_sifre_degistirme(self):
        adm = self.giris("admin", "admin")
        r = self.c.get("/api/panel", headers=adm)
        self.assertEqual(r.status_code, 403)
        self.assertIn("şifrenizi", r.json()["hata"])
        self.assertEqual(self.c.post("/api/sifre", headers=adm, json={"eski": "admin", "yeni": "123"}).status_code, 400)
        self.assertEqual(self.c.post("/api/sifre", headers=adm, json={"eski": "admin", "yeni": "YeniSifre1"}).status_code, 200)
        r = self.c.get("/api/panel", headers=adm)
        self.assertEqual(r.status_code, 200, r.text)
        self.assertIn("nakit", r.json()["ozet"])

    def test_uyum_yapilacaklar_ve_rol_rehberi_api(self):
        mudur = self.giris("ali")
        ben = self.c.get("/api/ben", headers=mudur).json()
        self.assertEqual(ben["sayfa_modulleri"]["Mevzuat Uyum Denetimi"], "uyum")

        uyum_ = self.c.get("/api/uyum", headers=mudur, params={"yil": YIL})
        self.assertEqual(uyum_.status_code, 200, uyum_.text)
        self.assertEqual(uyum_.json()["yil"], YIL)
        self.assertIn("puan", uyum_.json())
        self.assertIn("hata", uyum_.json()["sayilar"])
        self.assertTrue(all("ilgili" in x and "sorumlular" in x for x in uyum_.json()["bulgular"]))

        gorevler = self.c.get("/api/yapilacaklar", headers=mudur, params={"yil": YIL})
        self.assertEqual(gorevler.status_code, 200, gorevler.text)
        self.assertIn("bulgular", gorevler.json())
        self.assertIn("sureler", gorevler.json())

        kilavuz = self.c.get("/api/kilavuz", headers=mudur)
        self.assertEqual(kilavuz.status_code, 200, kilavuz.text)
        self.assertTrue(kilavuz.json()["rehberler"])
        self.assertIn("adimlar", kilavuz.json()["rehberler"][0])
        tumu = self.c.get("/api/kilavuz", headers=mudur, params={"tumu": "true"})
        self.assertGreaterEqual(len(tumu.json()["rehberler"]), len(kilavuz.json()["rehberler"]))
        yardim = self.c.get("/api/kilavuz/sayfa", headers=mudur, params={"ad": "Kasa / Banka"})
        self.assertEqual(yardim.status_code, 200, yardim.text)
        self.assertTrue(yardim.json()["nasil"])
        self.assertIn("dikkat", yardim.json())
        self.assertEqual(self.c.get("/api/kilavuz/sayfa", headers=mudur,
                                    params={"ad": "Bilinmeyen sayfa"}).status_code, 404)

        ogretmen = self.giris("deniz")
        denetim = self.c.get("/api/uyum", headers=ogretmen, params={"yil": YIL})
        self.assertEqual(denetim.status_code, 200, denetim.text)
        self.assertTrue(denetim.json()["bolum_kisitli"])
        self.assertEqual(self.c.get("/api/uyum", headers=ogretmen, params={"yil": 2300}).status_code, 422)

    def test_ogretmen_fazla_calisma_ve_mobil_onay(self):
        ogr = self.giris("deniz", cihaz="Samsung A54")
        sec = self.c.get(f"/api/fazla-calisma/secenekler?yil={YIL}&ay=3", headers=ogr).json()
        self.assertEqual([p["id"] for p in sec["personeller"]], [self.p_ogr])
        gun = next(g for g in sec["gunler"] if g["tur"] == "ogretim")
        istek = {"yil": YIL, "ay": 3, "personel_id": self.p_ogr2, "puantaj": {str(gun["gun"]): gun["azami"]},
                 "aciklama": "Zeytinyağı dolumu", "onizleme": True}
        on = self.c.post("/api/fazla-calisma", headers=ogr, json=istek).json()
        self.assertEqual(on["hatalar"], [])
        self.assertEqual(on["kayit"]["personel_id"], self.p_ogr)  # öğretmen yalnız kendisi için girer
        self.assertAlmostEqual(on["kayit"]["brut"], gun["azami"] * 67.5)
        istek["onizleme"] = False
        r = self.c.post("/api/fazla-calisma", headers=ogr, json=istek)
        self.assertEqual(r.status_code, 200, r.text)
        fid = r.json()["id"]
        fazla = dict(istek, puantaj={str(gun["gun"]): 99})
        r = self.c.post("/api/fazla-calisma", headers=ogr, json=fazla)
        self.assertEqual(r.status_code, 400)
        self.assertIn("günlük azami", r.json()["hata"])

        r = self.c.post(f"/api/onay/{FC}/{fid}/gonder", headers=ogr, json={"aciklama": "Mobilden gönderildi"})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(self.c.put(f"/api/fazla-calisma/{fid}", headers=ogr, json=istek).status_code, 400)

        sef, tmy, mdr = self.giris("can"), self.giris("veli"), self.giris("ali", cihaz="iPhone 15")
        self.assertEqual(self.c.post(f"/api/onay/{FC}/{fid}/karar", headers=mdr, json={"onay": True}).status_code, 400)
        for h in (sef, tmy, mdr):
            bekleyen = self.c.get("/api/onay", headers=h).json()["bana_bekleyen"]
            self.assertIn(fid, [b["belge_id"] for b in bekleyen if b["belge_turu"] == "Fazla Çalışma"])
            belge = self.c.get(f"/api/onay/{FC}/{fid}", headers=h).json()
            self.assertTrue(belge["islemler"]["karar"])
            self.assertEqual(belge["detay"]["ad_soyad"], "Deniz Öğretmen")
            r = self.c.post(f"/api/onay/{FC}/{fid}/karar", headers=h, json={"onay": True, "aciklama": "Uygundur"})
            self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(self.db.deger("SELECT durum FROM fazla_calisma WHERE id = ?", (fid,)), "Onaylandı")
        self.assertEqual(self.c.delete(f"/api/fazla-calisma/{fid}", headers=ogr).status_code, 400)

        kayit = self.c.get("/api/islem-kaydi?istemci=Mobil", headers=mdr).json()["kayitlar"]
        ekleme = next(k for k in kayit if k["tablo"] == "fazla_calisma" and k["islem"] == "Ekleme")
        self.assertEqual((ekleme["kim"], ekleme["rol"], ekleme["bilgisayar"], ekleme["istemci"]),
                         ("Deniz Öğretmen", M.ROL_OGRETMEN, "Samsung A54", "Mobil"))
        ayrinti = self.c.get(f"/api/islem-kaydi/{ekleme['id']}", headers=mdr).json()
        self.assertTrue(any(x["alan"] == "brut" for x in ayrinti["degisiklikler"]))
        gecmis = self.db.q("SELECT karar, rol FROM onay_kaydi WHERE belge_turu = 'Fazla Çalışma' AND belge_id = ? ORDER BY id", (fid,))
        self.assertEqual([g["rol"] for g in gecmis][-3:], [M.ROL_SEF, M.ROL_TMY, M.ROL_MUDUR])

    def test_is_talebi_ve_bolum_kisiti(self):
        ogr = self.giris("deniz")
        sec = self.c.get("/api/secenekler/siparis", headers=ogr).json()
        self.assertEqual([b["id"] for b in sec["bolumler"]], [self.bolum])
        talep = {"cari_id": self.musteri, "tarih": f"{YIL}-04-01", "konu": "Belediye için zeytinyağı",
                 "kalemler": [{"urun_id": self.urun, "miktar": 10}]}
        r = self.c.post("/api/siparisler", headers=ogr, json=talep)
        self.assertEqual(r.status_code, 200, r.text)
        s = self.c.get(f"/api/siparisler/{r.json()['id']}", headers=ogr).json()
        self.assertEqual((s["durum"], s["bolum_id"], s["sozlesme_bedeli"]), ("Teklif", self.bolum, 4000))
        self.assertEqual(s["malzeme_maliyeti"], 2200)  # (5 x 40 + 20) x 10
        self.assertAlmostEqual(s["iscilik_tutari"], s["tahsis"]["azami_iscilik"])
        self.assertEqual(s["kalemler"][0]["birim_fiyat"], 400)
        r = self.c.post("/api/siparisler", headers=ogr, json=dict(talep, bolum_id=self.bolum2))
        self.assertEqual(r.status_code, 400)
        self.assertIn("kendi bölümünüz", r.json()["hata"])
        r = self.c.post("/api/siparisler", headers=ogr, json=dict(talep, kalemler=[{"urun_id": self.urun, "miktar": 1,
                                                                                    "birim_fiyat": 300}]))
        self.assertTrue(any("tespit edilen satış fiyatının" in u for u in r.json()["uyarilar"]))
        self.assertEqual(self.c.post("/api/siparisler", headers=ogr, json={"cari_id": self.musteri}).status_code, 400)
        diger = self.db.ekle("siparis", {"siparis_no": "SIP-X", "tarih": f"{YIL}-04-02", "bolum_id": self.bolum2,
                                         "konu": "Bilişim işi", "durum": "Teklif"})
        liste = self.c.get(f"/api/siparisler?yil={YIL}", headers=ogr).json()["siparisler"]
        self.assertNotIn(diger, [x["id"] for x in liste])
        self.assertEqual(self.c.get(f"/api/siparisler/{diger}", headers=ogr).status_code, 403)

    def test_is_talebi_duzenleme_cakisma_iliski_engeli_ve_silme(self):
        ogr = self.giris("deniz")
        talep = {"cari_id": self.musteri, "tarih": f"{YIL}-04-03", "konu": "Düzenleme testi",
                 "kalemler": [{"urun_id": self.urun, "miktar": 2}]}
        ekle = self.c.post("/api/siparisler", headers=ogr, json=talep)
        self.assertEqual(ekle.status_code, 200, ekle.text)
        sid = ekle.json()["id"]
        self.db.guncelle("siparis", sid, {"pesinat": 50, "sozlesme_suresi_ay": 12, "bakanlik_izni": 1})
        detay = self.c.get(f"/api/siparisler/{sid}", headers=ogr).json()
        self.assertTrue(detay["duzenlenebilir"])
        self.assertTrue(detay["silinebilir"])
        izleyici = self.c.get(f"/api/siparisler/{sid}", headers=self.giris("deneme")).json()
        self.assertFalse(izleyici["duzenlenebilir"])
        self.assertFalse(izleyici["silinebilir"])

        masaustu = Veritabani(self.yol)
        masaustu.oturum_ac({"kullanici_adi": "veli", "ad_soyad": "Veli Yardımcı", "rol": M.ROL_TMY},
                           bilgisayar="TMY-PC")
        masaustu.guncelle("siparis", sid, {"aciklama": "Masaüstünde değişti"})
        masaustu.kapat()
        govde = {**talep, "konu": "Güncel iş talebi", "kalemler": [{"urun_id": self.urun, "miktar": 4}],
                 "surum": detay["surum"]}
        cakisma = self.c.put(f"/api/siparisler/{sid}", headers=ogr, json=govde)
        self.assertEqual(cakisma.status_code, 409, cakisma.text)
        self.assertIn("Veli Yardımcı", cakisma.json()["hata"])

        guncel_surumu = self.c.get(f"/api/siparisler/{sid}", headers=ogr).json()["surum"]
        govde["surum"] = guncel_surumu
        guncelle = self.c.put(f"/api/siparisler/{sid}", headers=ogr, json=govde)
        self.assertEqual(guncelle.status_code, 200, guncelle.text)
        self.assertEqual(guncelle.json()["siparis_no"], detay["siparis_no"])
        saklanan = self.db.q1("SELECT tarih, konu, durum, sozlesme_bedeli, pesinat, sozlesme_suresi_ay, bakanlik_izni "
                              "FROM siparis WHERE id = ?", (sid,))
        self.assertEqual((saklanan["tarih"], saklanan["konu"], saklanan["durum"], saklanan["sozlesme_bedeli"],
                          saklanan["pesinat"], saklanan["sozlesme_suresi_ay"], saklanan["bakanlik_izni"]),
                         (f"{YIL}-04-03", "Güncel iş talebi", "Teklif", 1600, 50, 12, 1))
        self.assertEqual(self.db.deger("SELECT miktar FROM siparis_kalem WHERE siparis_id = ?", (sid,)), 4)

        for tur in ("kesif", "proforma", "teklif", "sozlesme", "maliyet", "malzeme",
                    "gecici", "istek", "gorev", "kabul"):
            belge = self.c.get(f"/api/siparisler/{sid}/belge/{tur}", headers=ogr)
            self.assertEqual(belge.status_code, 200, f"{tur}: {belge.text}")
            self.assertTrue(base64.b64decode(belge.json()["icerik"]).startswith(b"%PDF"), tur)
        proforma = self.db.q1("SELECT proforma_no, proforma_tarih FROM siparis WHERE id = ?", (sid,))
        self.assertTrue(proforma["proforma_no"])
        self.assertTrue(proforma["proforma_tarih"])
        self.assertEqual(self.c.get(f"/api/siparisler/{sid}/belge/hatalı", headers=ogr).status_code, 404)

        self.db.ekle("fatura", {"tur": "Satış", "siparis_id": sid})
        detay = self.c.get(f"/api/siparisler/{sid}", headers=ogr).json()
        self.assertFalse(detay["silinebilir"])
        engel = self.c.delete(f"/api/siparisler/{sid}", headers=ogr, params={"surum": detay["surum"]})
        self.assertEqual(engel.status_code, 400, engel.text)
        self.db.ex("DELETE FROM fatura WHERE siparis_id = ?", (sid,))
        self.db.ekle("siparis_is_plani", {"siparis_id": sid, "kalem": "Kontrol"})
        self.db.ekle("siparis_gorevli", {"siparis_id": sid, "kisi_turu": "Personel", "kisi_id": self.p_ogr})
        detay = self.c.get(f"/api/siparisler/{sid}", headers=ogr).json()
        sil = self.c.delete(f"/api/siparisler/{sid}", headers=ogr, params={"surum": detay["surum"]})
        self.assertEqual(sil.status_code, 200, sil.text)
        self.assertIsNone(self.db.q1("SELECT id FROM siparis WHERE id = ?", (sid,)))
        for tablo in ("siparis_kalem", "siparis_is_plani", "siparis_gorevli"):
            self.assertEqual(self.db.deger(f"SELECT COUNT(*) FROM {tablo} WHERE siparis_id = ?", (sid,)), 0)

    def test_siparis_faturalama_stok_cikisi_ve_excel_disari_aktarimi(self):
        talep = {"cari_id": self.musteri, "tarih": f"{YIL}-04-05", "konu": "Sipariş yan etkileri",
                 "kalemler": [{"urun_id": self.urun, "miktar": 3}]}
        ekle = self.c.post("/api/siparisler", headers=self.giris("deniz"), json=talep)
        self.assertEqual(ekle.status_code, 200, ekle.text)
        sid = ekle.json()["id"]
        yetkili = self.giris("ayse")
        onceki_akis = onay.akis_acik()
        try:
            onay.akis_ayarla(False)
            detay = self.c.get(f"/api/siparisler/{sid}", headers=yetkili).json()
            fatura = self.c.post(f"/api/siparisler/{sid}/faturala", headers=yetkili, json={"surum": detay["surum"]})
            self.assertEqual(fatura.status_code, 200, fatura.text)
            fatura_id = fatura.json()["id"]
            self.assertEqual((fatura.json()["matrah"], fatura.json()["genel_toplam"]), (1200, 1440))
            self.assertEqual(self.db.deger("SELECT COUNT(*) FROM fatura_kalem WHERE fatura_id = ?", (fatura_id,)), 1)
            kalan = self.c.post(f"/api/siparisler/{sid}/faturala", headers=yetkili,
                                json={"surum": self.c.get(f"/api/siparisler/{sid}", headers=yetkili).json()["surum"]})
            self.assertEqual(kalan.status_code, 400, kalan.text)
        finally:
            onay.akis_ayarla(onceki_akis)

        malzeme_id = self.db.ekle("malzeme", {"ad": "API çıkış malzemesi", "tur": "İlk Madde ve Malzeme", "birim": "kg"})
        self.db.ekle("stok_hareket", {"tarih": f"{YIL}-04-01", "malzeme_id": malzeme_id, "yon": "Giriş",
                                      "miktar": 5, "birim_fiyat": 12, "kaynak": "Manuel"})
        stok_detay = self.c.post(
            f"/api/siparisler/{sid}/stok-cikisi", headers=yetkili,
            json={"malzeme_id": malzeme_id, "miktar": 6, "tarih": f"{YIL}-04-06", "neden": "Üretime Sevk"})
        self.assertEqual(stok_detay.status_code, 400, stok_detay.text)
        stok = self.c.post(
            f"/api/siparisler/{sid}/stok-cikisi", headers=yetkili,
            json={"malzeme_id": malzeme_id, "miktar": 2, "tarih": f"{YIL}-04-06",
                  "neden": "Üretime Sevk", "aciklama": "Sipariş için"})
        self.assertEqual(stok.status_code, 200, stok.text)
        self.assertEqual(stok.json()["birim_fiyat"], 12)
        self.assertEqual(self.db.deger("SELECT siparis_id FROM stok_hareket WHERE id = ?", (stok.json()["id"],)), sid)
        self.assertEqual(self.db.deger("SELECT COUNT(*) FROM yevmiye WHERE kaynak = 'Stok' AND kaynak_id = ?",
                                       (stok.json()["id"],)), 1)

        disari = self.c.get(f"/api/siparisler/disa-aktar?yil={YIL}&durum=tumu", headers=yetkili)
        self.assertEqual(disari.status_code, 200, disari.text)
        self.assertTrue(base64.b64decode(disari.json()["icerik"]).startswith(b"PK"))

    def test_fatura_olusturma_duzenleme_odeme_iptal_ve_excel_api(self):
        sayman = self.giris("ayse")
        def temizle_fatura_testi():
            no_listesi = ("FLET-2026-001", "FLET-2026-002", "FLET-2026-003", "ALIŞ-API-001")
            faturalar = self.db.q("SELECT id FROM fatura WHERE belge_no IN (?, ?, ?, ?)", no_listesi)
            ids = [f["id"] for f in faturalar]
            if not ids:
                return
            yer_tutucular = ",".join("?" for _ in ids)
            with self.db.islem():
                for hareket in self.db.q(f"SELECT id FROM kasa_banka WHERE fatura_id IN ({yer_tutucular})", ids):
                    for fis in self.db.q("SELECT id FROM yevmiye WHERE kaynak = 'KasaBanka' AND kaynak_id = ?",
                                         (hareket["id"],)):
                        self.db.ex("DELETE FROM yevmiye_satir WHERE yevmiye_id = ?", (fis["id"],))
                        self.db.sil("yevmiye", fis["id"])
                    self.db.sil("kasa_banka", hareket["id"])
                for kaynak in ("Fatura",):
                    for fis in self.db.q(f"SELECT id FROM yevmiye WHERE kaynak = ? AND kaynak_id IN ({yer_tutucular})",
                                         (kaynak, *ids)):
                        self.db.ex("DELETE FROM yevmiye_satir WHERE yevmiye_id = ?", (fis["id"],))
                        self.db.sil("yevmiye", fis["id"])
                self.db.ex(f"DELETE FROM stok_hareket WHERE kaynak = 'Fatura' AND kaynak_id IN ({yer_tutucular})", ids)
                self.db.ex(f"DELETE FROM fatura_kalem WHERE fatura_id IN ({yer_tutucular})", ids)
                self.db.ex(f"DELETE FROM fatura WHERE id IN ({yer_tutucular})", ids)
                self.db.ex("DELETE FROM malzeme WHERE ad = 'Alış fatura malzemesi'")
                self.db.ex("DELETE FROM cari WHERE unvan = 'API Tedarikçisi'")
        self.addCleanup(temizle_fatura_testi)
        govde = {
            "belge_turu": "e-Arşiv Fatura", "belge_no": "FLET-2026-001", "tarih": f"{YIL}-04-10",
            "cari_id": self.musteri, "tevkifat": "Yok", "kalemler": [{
                "urun_id": self.urun, "aciklama": "Test ürünü", "miktar": 2, "birim": "Adet",
                "birim_fiyat": 10, "iskonto_orani": 0.1, "kdv_orani": 0.2,
            }],
        }
        ekle = self.c.post("/api/faturalar?tur=Satış", headers=sayman, json=govde)
        self.assertEqual(ekle.status_code, 200, ekle.text)
        fid = ekle.json()["id"]
        self.assertEqual((ekle.json()["matrah"], ekle.json()["kdv"], ekle.json()["genel_toplam"]), (18, 3.6, 21.6))
        detay = self.c.get(f"/api/faturalar/{fid}", headers=sayman)
        self.assertEqual(detay.status_code, 200, detay.text)
        self.assertEqual(detay.json()["kalemler"][0]["iskonto"], 2)

        duzenle_govde = {**govde, "surum": detay.json()["surum"], "belge_no": "FLET-2026-002",
                         "kalemler": [{**govde["kalemler"][0], "miktar": 3}]}
        duzenle = self.c.put(f"/api/faturalar/{fid}", headers=sayman, json=duzenle_govde)
        self.assertEqual(duzenle.status_code, 200, duzenle.text)
        self.assertEqual(duzenle.json()["matrah"], 27)

        odeme = self.c.post(f"/api/faturalar/{fid}/odeme", headers=sayman,
                            json={"tarih": f"{YIL}-04-11", "hesap": "Banka (102)", "tutar": 5})
        self.assertEqual(odeme.status_code, 200, odeme.text)
        self.assertEqual(odeme.json()["kalan"], 27 * 1.2 - 5)
        self.assertAlmostEqual(self.db.deger("SELECT hasilat_tutari FROM kasa_banka WHERE fatura_id = ?", (fid,)),
                               5 * 27 / (27 * 1.2), places=2)
        guncelleme_engeli = self.c.put(f"/api/faturalar/{fid}", headers=sayman, json=duzenle_govde)
        self.assertEqual(guncelleme_engeli.status_code, 400, guncelleme_engeli.text)
        iptal_engeli = self.c.delete(f"/api/faturalar/{fid}", headers=sayman, params={"surum": duzenle.json()["surum"]})
        self.assertEqual(iptal_engeli.status_code, 400, iptal_engeli.text)

        disari = self.c.get(f"/api/faturalar/disa-aktar?tur=Satış&yil={YIL}", headers=sayman)
        self.assertEqual(disari.status_code, 200, disari.text)
        self.assertTrue(base64.b64decode(disari.json()["icerik"]).startswith(b"PK"))

        iptal_govde = {**govde, "belge_no": "FLET-2026-003"}
        iptal_fatura = self.c.post("/api/faturalar?tur=Satış", headers=sayman, json=iptal_govde)
        self.assertEqual(iptal_fatura.status_code, 200, iptal_fatura.text)
        iptal_detay = self.c.get(f"/api/faturalar/{iptal_fatura.json()['id']}", headers=sayman).json()
        iptal = self.c.delete(f"/api/faturalar/{iptal_fatura.json()['id']}", headers=sayman,
                              params={"surum": iptal_detay["surum"]})
        self.assertEqual(iptal.status_code, 200, iptal.text)
        self.assertTrue(self.c.get(f"/api/faturalar/{iptal_fatura.json()['id']}", headers=sayman).json()["iptal"])

        tedarikci = self.db.ekle("cari", {"unvan": "API Tedarikçisi", "tur": "Tedarikçi"})
        malzeme_id = self.db.ekle("malzeme", {"ad": "Alış fatura malzemesi", "tur": "İlk Madde ve Malzeme",
                                              "birim": "kg"})
        alis = self.c.post("/api/faturalar?tur=Alış", headers=sayman, json={
            "belge_turu": "e-Fatura", "belge_no": "ALIŞ-API-001", "tarih": f"{YIL}-04-12",
            "cari_id": tedarikci, "gider_turu": "İlk Madde ve Malzeme (Stoğa)",
            "temin_usulu": "Doğrudan Temin", "kalemler": [{
                "malzeme_id": malzeme_id, "aciklama": "Stok alımı", "miktar": 2, "birim": "kg",
                "birim_fiyat": 25, "kdv_orani": 0.2,
            }],
        })
        self.assertEqual(alis.status_code, 200, alis.text)
        hareket = self.db.q1("SELECT * FROM stok_hareket WHERE kaynak = 'Fatura' AND kaynak_id = ?",
                             (alis.json()["id"],))
        self.assertEqual((hareket["yon"], hareket["miktar"], hareket["birim_fiyat"]), ("Giriş", 2, 25))
        self.assertEqual(self.c.get(f"/api/faturalar/disa-aktar?tur=Alış&yil={YIL}", headers=sayman).status_code, 200)
        onbilgi = self.c.get(f"/api/faturalar/{alis.json()['id']}/dogrudan-temin", headers=sayman)
        self.assertEqual(onbilgi.status_code, 200, onbilgi.text)
        paket = self.c.post("/api/belgeler/dogrudan-temin", headers=sayman, json={
            **onbilgi.json(), "konu": "Üretim için stok alımı",
            "firmalar": [{"ad": "Birinci Firma", "tutar": 50},
                                            {"ad": "İkinci Firma", "tutar": 55},
                                            {"ad": "Üçüncü Firma", "tutar": 60}],
            "gorevliler": ["Ali Yetkili"],
        })
        self.assertEqual(paket.status_code, 200, paket.text)
        self.assertTrue(base64.b64decode(paket.json()["icerik"]).startswith(b"%PDF"))
        self.assertFalse(paket.json()["uyarilar"])

    def test_kasa_api_manuel_hareket_duzenleme_silme_bakiye_yetkisi_ve_excel(self):
        sayman = self.giris("ayse")
        hareket_ids = []
        kaynak_id = self.db.ekle("kasa_banka", {
            "tarih": f"{YIL}-06-01", "hesap": "Banka (102)", "yon": "Giriş",
            "tur": "Tahsilat", "tutar": 20, "hasilat_tutari": 20,
            "kaynak": "HazinePayi", "kaynak_id": 12345,
        })

        def temizle():
            with self.db.islem():
                for hareket_id in [*hareket_ids, kaynak_id]:
                    for fis in self.db.q("SELECT id FROM yevmiye WHERE kaynak = 'KasaBanka' AND kaynak_id = ?",
                                         (hareket_id,)):
                        self.db.ex("DELETE FROM yevmiye_satir WHERE yevmiye_id = ?", (fis["id"],))
                        self.db.sil("yevmiye", fis["id"])
                    self.db.ex("DELETE FROM kasa_banka WHERE id = ?", (hareket_id,))
        self.addCleanup(temizle)

        secenekler = self.c.get("/api/kasa", headers=sayman, params={"yil": YIL})
        self.assertEqual(secenekler.status_code, 200, secenekler.text)
        self.assertIn("hesaplar", secenekler.json())
        self.assertIn("faturalar", secenekler.json())

        govde = {"tarih": f"{YIL}-06-02", "hesap": "Banka (102)", "yon": "Giriş",
                 "tur": "Faiz/Diğer Gelir", "tutar": 100, "aciklama": "Kasa API testi"}
        ekle = self.c.post("/api/kasa", headers=sayman, json=govde)
        self.assertEqual(ekle.status_code, 200, ekle.text)
        hareket_id = ekle.json()["id"]
        hareket_ids.append(hareket_id)
        self.assertEqual(ekle.json()["hasilat_tutari"], 100)
        detay = self.c.get(f"/api/kasa/{hareket_id}", headers=sayman)
        self.assertEqual(detay.status_code, 200, detay.text)
        self.assertTrue(detay.json()["duzenlenebilir"])

        guncel = {**govde, "tutar": 120, "hasilat_tutari": 110, "surum": detay.json()["surum"]}
        duzenle = self.c.put(f"/api/kasa/{hareket_id}", headers=sayman, json=guncel)
        self.assertEqual(duzenle.status_code, 200, duzenle.text)
        self.assertEqual(duzenle.json()["hasilat_tutari"], 110)
        self.assertEqual(self.db.deger("SELECT COUNT(*) FROM yevmiye WHERE kaynak = 'KasaBanka' AND kaynak_id = ?",
                                       (hareket_id,)), 1)
        cakisma = self.c.put(f"/api/kasa/{hareket_id}", headers=self.giris("ayse", cihaz="İkinci cihaz"),
                             json=guncel)
        self.assertEqual(cakisma.status_code, 409, cakisma.text)

        yetersiz = self.c.post("/api/kasa", headers=sayman, json={
            "tarih": f"{YIL}-06-03", "hesap": "Kasa (100)", "yon": "Çıkış",
            "tur": "Vergi Ödemesi", "tutar": 1000,
        })
        self.assertEqual(yetersiz.status_code, 409, yetersiz.text)
        onayli = self.c.post("/api/kasa", headers=sayman, json={
            "tarih": f"{YIL}-06-03", "hesap": "Kasa (100)", "yon": "Çıkış",
            "tur": "Vergi Ödemesi", "tutar": 1000, "bakiye_uyarisi_onaylandi": True,
        })
        self.assertEqual(onayli.status_code, 200, onayli.text)
        hareket_ids.append(onayli.json()["id"])

        kaynak_detay = self.c.get(f"/api/kasa/{kaynak_id}", headers=sayman)
        self.assertEqual(kaynak_detay.status_code, 200, kaynak_detay.text)
        self.assertFalse(kaynak_detay.json()["duzenlenebilir"])
        self.assertEqual(self.c.delete(f"/api/kasa/{kaynak_id}", headers=sayman,
                                       params={"surum": kaynak_detay.json()["surum"]}).status_code, 400)
        kaynak_duzenle = self.c.put(f"/api/kasa/{kaynak_id}", headers=sayman, json={
            "tarih": f"{YIL}-06-01", "hesap": "Banka (102)", "yon": "Giriş",
            "tur": "Tahsilat", "tutar": 20, "surum": kaynak_detay.json()["surum"],
        })
        self.assertEqual(kaynak_duzenle.status_code, 400, kaynak_duzenle.text)
        salt_okunur = self.c.post("/api/kasa", headers=self.giris("deneme"), json=govde)
        self.assertEqual(salt_okunur.status_code, 403, salt_okunur.text)
        excel = self.c.get(f"/api/kasa/disa-aktar?yil={YIL}&ay=6", headers=sayman)
        self.assertEqual(excel.status_code, 200, excel.text)
        self.assertTrue(base64.b64decode(excel.json()["icerik"]).startswith(b"PK"))

        sil = self.c.delete(f"/api/kasa/{hareket_id}", headers=sayman, params={"surum": duzenle.json()["surum"]})
        self.assertEqual(sil.status_code, 200, sil.text)
        hareket_ids.remove(hareket_id)

    def test_stok_api_hareket_maliyet_yetersiz_stok_silme_ve_excel(self):
        sayman = self.giris("ayse")
        malzeme_id = self.db.ekle("malzeme", {"ad": "Stok API hareket testi", "tur": "İlk Madde ve Malzeme",
                                              "birim": "kg"})
        hareket_ids = []
        fatura_kaynak_id = self.db.ekle("stok_hareket", {
            "tarih": f"{YIL}-07-01", "malzeme_id": malzeme_id, "yon": "Giriş", "miktar": 1,
            "birim_fiyat": 10, "kaynak": "Fatura", "kaynak_id": 98765,
        })
        hareket_ids.append(fatura_kaynak_id)

        def temizle():
            with self.db.islem():
                for hareket_id in hareket_ids:
                    for fis in self.db.q("SELECT id FROM yevmiye WHERE kaynak = 'Stok' AND kaynak_id = ?",
                                         (hareket_id,)):
                        self.db.ex("DELETE FROM yevmiye_satir WHERE yevmiye_id = ?", (fis["id"],))
                        self.db.sil("yevmiye", fis["id"])
                    self.db.ex("DELETE FROM stok_hareket WHERE id = ?", (hareket_id,))
                self.db.ex("DELETE FROM malzeme WHERE id = ?", (malzeme_id,))
        self.addCleanup(temizle)

        liste = self.c.get(f"/api/stok?yil={YIL}", headers=sayman)
        self.assertEqual(liste.status_code, 200, liste.text)
        self.assertIn("Stok API hareket testi", [m["ad"] for m in liste.json()["malzemeler"]])
        faturali = next(h for h in liste.json()["hareketler"] if h["id"] == fatura_kaynak_id)
        self.assertFalse(faturali["silinebilir"])

        giris = {"tarih": f"{YIL}-07-02", "malzeme_id": malzeme_id, "yon": "Giriş",
                 "neden": "Devir / Açılış", "miktar": 10, "birim_fiyat": 20}
        ekle = self.c.post("/api/stok/hareketler", headers=sayman, json=giris)
        self.assertEqual(ekle.status_code, 200, ekle.text)
        giris_id = ekle.json()["id"]
        hareket_ids.append(giris_id)
        self.assertEqual(ekle.json()["tutar"], 200)

        sifir_fiyat = self.c.post("/api/stok/hareketler", headers=sayman, json={
            **giris, "tarih": f"{YIL}-07-03", "miktar": 1, "birim_fiyat": 0,
        })
        self.assertEqual(sifir_fiyat.status_code, 400, sifir_fiyat.text)
        sifir_onayli = self.c.post("/api/stok/hareketler", headers=sayman, json={
            **giris, "tarih": f"{YIL}-07-03", "miktar": 1, "birim_fiyat": 0,
            "sifir_fiyat_onaylandi": True,
        })
        self.assertEqual(sifir_onayli.status_code, 200, sifir_onayli.text)
        hareket_ids.append(sifir_onayli.json()["id"])

        cikis = {"tarih": f"{YIL}-07-04", "malzeme_id": malzeme_id, "yon": "Çıkış",
                 "neden": "Üretime Sevk", "miktar": 3}
        cikis_ekle = self.c.post("/api/stok/hareketler", headers=sayman, json=cikis)
        self.assertEqual(cikis_ekle.status_code, 200, cikis_ekle.text)
        cikis_id = cikis_ekle.json()["id"]
        hareket_ids.append(cikis_id)
        self.assertAlmostEqual(cikis_ekle.json()["birim_fiyat"], 17.5)
        self.assertEqual(self.db.deger("SELECT COUNT(*) FROM yevmiye WHERE kaynak = 'Stok' AND kaynak_id = ?",
                                       (cikis_id,)), 1)

        yetersiz = self.c.post("/api/stok/hareketler", headers=sayman,
                               json={**cikis, "miktar": 10})
        self.assertEqual(yetersiz.status_code, 400, yetersiz.text)
        sil = self.c.delete(f"/api/stok/hareketler/{cikis_id}", headers=sayman,
                            params={"surum": cikis_ekle.json()["surum"]})
        self.assertEqual(sil.status_code, 200, sil.text)
        self.assertEqual(self.db.deger("SELECT COUNT(*) FROM yevmiye WHERE kaynak = 'Stok' AND kaynak_id = ?",
                                       (cikis_id,)), 0)
        hareket_ids.remove(cikis_id)

        faturali_sil = self.c.delete(f"/api/stok/hareketler/{fatura_kaynak_id}", headers=sayman,
                                     params={"surum": liste.json()["surum"]})
        self.assertEqual(faturali_sil.status_code, 400, faturali_sil.text)
        excel = self.c.get(f"/api/stok/disa-aktar?yil={YIL}", headers=sayman)
        self.assertEqual(excel.status_code, 200, excel.text)
        self.assertTrue(base64.b64decode(excel.json()["icerik"]).startswith(b"PK"))

    def test_personel_ogrenci_excel_onizleme_dogrulama_ve_ice_aktarma(self):
        from openpyxl import Workbook

        sayman = self.giris("ayse")
        yonetici = self.giris("sistem")

        def xlsx_icerigi(basliklar, satirlar):
            wb = Workbook()
            ws = wb.active
            ws.title = "Aktarım"
            ws.append(basliklar)
            for satir in satirlar:
                ws.append(satir)
            akim = BytesIO()
            wb.save(akim)
            return base64.b64encode(akim.getvalue()).decode("ascii")

        govdeler = {
            "personel": {"dosya_adi": "personel.xlsx", "icerik": xlsx_icerigi(
                ["Sicil No", "T.C. Kimlik No", "Adı Soyadı", "Unvan", "IBAN", "Telefon"], [
                    ["P-01", "10000000146", "Excel Personeli", "Öğretmen", "", "05550000001"],
                    ["P-02", "10000000146", "Mükerrer Personel", "Öğretmen", "", ""],
                    ["P-03", "", "İkinci Excel Personeli", "Öğretmen", "", ""],
                    ["O-01", "", "Excel Öğrencisi", "Öğrenci", "", ""],
                ])},
            "ogrenci": {"dosya_adi": "ogrenciler.xlsx", "icerik": xlsx_icerigi(
                ["Okul No", "T.C. Kimlik No", "Adı Soyadı", "Sınıf", "Görevi", "IBAN"], [
                    ["O-02", "", "İçe Aktarılan Öğrenci", "10-A", "Öğrenci", "geçersiz"],
                    ["P-04", "", "Personel Satırı", "", "Öğretmen", ""],
                ])},
        }

        def temizle():
            self.db.ex("DELETE FROM personel WHERE ad_soyad IN (?, ?)",
                       ("Excel Personeli", "İkinci Excel Personeli"))
            self.db.ex("DELETE FROM ogrenci WHERE ad_soyad = ?", ("İçe Aktarılan Öğrenci",))
        self.addCleanup(temizle)

        onizleme = self.c.post("/api/tanimlar/personel/excel-onizle", headers=sayman, json=govdeler["personel"])
        self.assertEqual(onizleme.status_code, 200, onizleme.text)
        ozet = onizleme.json()["ozet"]
        self.assertEqual((ozet["aktarilacak"], ozet["mevcut"], ozet["elenen"]), (2, 1, 1))
        self.assertEqual(onizleme.json()["sayfa"], "Aktarım")
        aktar = self.c.post("/api/tanimlar/personel/excel-aktar", headers=sayman, json=govdeler["personel"])
        self.assertEqual(aktar.status_code, 200, aktar.text)
        self.assertEqual(aktar.json()["eklenen"], 2)
        personel = self.db.q1("SELECT * FROM personel WHERE ad_soyad = 'Excel Personeli'")
        self.assertEqual((personel["sicil_no"], personel["dose_gorevi"], personel["aktif"]),
                         ("P-01", "Öğretmen", 1))

        ogrenci_onizleme = self.c.post("/api/tanimlar/ogrenci/excel-onizle",
                                       headers=yonetici, json=govdeler["ogrenci"])
        self.assertEqual(ogrenci_onizleme.status_code, 200, ogrenci_onizleme.text)
        self.assertEqual((ogrenci_onizleme.json()["ozet"]["aktarilacak"],
                          ogrenci_onizleme.json()["ozet"]["elenen"]), (1, 1))
        self.assertTrue(any("geçersiz IBAN" in n for n in ogrenci_onizleme.json()["ozet"]["notlar"]))
        ogrenci_aktar = self.c.post("/api/tanimlar/ogrenci/excel-aktar", headers=yonetici,
                                    json=govdeler["ogrenci"])
        self.assertEqual(ogrenci_aktar.status_code, 200, ogrenci_aktar.text)
        self.assertEqual(ogrenci_aktar.json()["eklenen"], 1)
        ogrenci = self.db.q1("SELECT * FROM ogrenci WHERE ad_soyad = 'İçe Aktarılan Öğrenci'")
        self.assertEqual((ogrenci["okul_no"], ogrenci["aktif"], ogrenci["veli_muvafakat"], ogrenci["iban"]),
                         ("O-02", 1, 0, None))

        hatali_dosya = self.c.post("/api/tanimlar/personel/excel-onizle", headers=sayman,
                                   json={"dosya_adi": "hatali.xlsx", "icerik": "not-base64"})
        self.assertEqual(hatali_dosya.status_code, 400, hatali_dosya.text)
        salt_okunur = self.c.post("/api/tanimlar/personel/excel-aktar", headers=self.giris("deneme"),
                                  json=govdeler["personel"])
        self.assertEqual(salt_okunur.status_code, 403, salt_okunur.text)

    def test_masaustu_ile_cakisma(self):
        sef = self.giris("can")
        istek = {"yil": YIL, "ay": 4, "personel_id": self.p_ogr2, "saat": 2}
        fid = self.c.post("/api/fazla-calisma", headers=sef, json=istek).json()["id"]
        surum = self.c.get(f"/api/fazla-calisma?yil={YIL}&ay=4", headers=sef).json()["surum"]
        masaustu = Veritabani(self.yol)
        masaustu.oturum_ac({"kullanici_adi": "veli", "ad_soyad": "Veli Yardımcı", "rol": M.ROL_TMY}, bilgisayar="TMY-PC")
        masaustu.guncelle("fazla_calisma", fid, {"aciklama": "Masaüstünde düzeltildi"})
        masaustu.kapat()
        r = self.c.put(f"/api/fazla-calisma/{fid}", headers=sef, json=dict(istek, saat=3, surum=surum))
        self.assertEqual(r.status_code, 409)
        self.assertIn("Veli Yardımcı", r.json()["hata"])
        yeni = self.c.get(f"/api/fazla-calisma?yil={YIL}&ay=4", headers=sef).json()["surum"]
        r = self.c.put(f"/api/fazla-calisma/{fid}", headers=sef, json=dict(istek, saat=3, surum=yeni))
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(self.db.deger("SELECT saat FROM fazla_calisma WHERE id = ?", (fid,)), 3)
        r = self.c.post("/api/fazla-calisma", headers=sef, json=dict(istek, personel_id=self.p_bilisim))
        self.assertEqual(r.status_code, 403)

    def test_salt_okunur_lisans(self):
        ogr = self.giris("deniz")
        self.db.salt_okunur = True
        try:
            r = self.c.post("/api/fazla-calisma", headers=ogr, json={"yil": YIL, "ay": 5, "saat": 1})
            self.assertEqual(r.status_code, 403)
            self.assertEqual(self.c.get("/api/ben", headers=ogr).json()["moduller"]["fazla"], "g")
            self.assertEqual(self.c.get(f"/api/fazla-calisma?yil={YIL}&ay=5", headers=ogr).status_code, 200)
        finally:
            self.db.salt_okunur = False

    def test_islem_kaydi_filtreleri_tarih_araligini_kontrol_eder_ve_excel_sunar(self):
        import base64
        from openpyxl import load_workbook

        admin = self.giris("sistem")
        ogretmen = self.giris("deniz")
        bugun = date.today().isoformat()
        filtre = {"kullanici_adi": "sistem", "istemci": "Mobil",
                  "baslangic": bugun, "bitis": bugun}
        liste = self.c.get("/api/islem-kaydi", headers=admin, params=filtre)
        self.assertEqual(liste.status_code, 200, liste.text)
        self.assertTrue(liste.json()["kayitlar"])
        self.assertTrue(all(x["kullanici"] == "sistem" for x in liste.json()["kayitlar"]))
        self.assertEqual(self.c.get("/api/islem-kaydi/excel", headers=ogretmen).status_code, 403)

        disari = self.c.get("/api/islem-kaydi/excel", headers=admin, params=filtre)
        self.assertEqual(disari.status_code, 200, disari.text)
        self.assertEqual(disari.json()["dosya_adi"], "Islem_Kaydi.xlsx")
        icerik = base64.b64decode(disari.json()["icerik"])
        self.assertTrue(icerik.startswith(b"PK"))
        kitap = load_workbook(BytesIO(icerik), read_only=True)
        self.assertEqual(kitap["İşlem Kaydı"]["A3"].value, "Zaman")
        kitap.close()

        gecersiz = self.c.get("/api/islem-kaydi", headers=admin, params={"baslangic": "2026-99-99"})
        self.assertEqual(gecersiz.status_code, 400)
        ters = self.c.get("/api/islem-kaydi", headers=admin,
                          params={"baslangic": "2026-12-31", "bitis": "2026-01-01"})
        self.assertEqual(ters.status_code, 400)

    def test_yedekleme_api_yedek_olusturma_indirme_ve_geri_yukleme_korumasini_uygular(self):
        import base64

        sistem = self.giris("sistem")
        ogretmen = self.giris("deniz")
        once = {p.name for p in self.yedekler.glob("*.db")}
        yeni_dosyalar: set[str] = set()
        try:
            liste = self.c.get("/api/yedekler", headers=sistem)
            self.assertEqual(liste.status_code, 200, liste.text)
            self.assertTrue(liste.json()["yonetici"])
            self.assertEqual(self.c.get("/api/yedekler", headers=ogretmen).status_code, 403)

            olustur = self.c.post("/api/yedekler", headers=sistem)
            self.assertEqual(olustur.status_code, 200, olustur.text)
            ad = olustur.json()["dosya_adi"]
            yeni_dosyalar.add(ad)
            self.assertTrue((self.yedekler / ad).exists())

            indir = self.c.get(f"/api/yedekler/{ad}", headers=sistem)
            self.assertEqual(indir.status_code, 200, indir.text)
            icerik = base64.b64decode(indir.json()["icerik"])
            self.assertGreater(len(icerik), 1024)

            self.assertEqual(self.c.post("/api/yedekler/geri-yukle", headers=ogretmen,
                                         json={"dosya_adi": ad, "onay_metni": "GERI YUKLE"}).status_code, 403)
            self.assertEqual(self.c.post("/api/yedekler/geri-yukle", headers=sistem,
                                         json={"dosya_adi": ad, "onay_metni": "yanlis"}).status_code, 400)
            geri = self.c.post("/api/yedekler/geri-yukle", headers=sistem,
                               json={"dosya_adi": ad, "onay_metni": "GERI YUKLE"})
            self.assertEqual(geri.status_code, 200, geri.text)
            sonra = {p.name for p in self.yedekler.glob("*.db")}
            yeni_dosyalar |= (sonra - once)
        finally:
            for ad in yeni_dosyalar:
                (self.yedekler / ad).unlink(missing_ok=True)

    def test_yedekleme_api_yerel_db_yuklemesini_dogrular_ve_yalniz_yoneticilere_acar(self):
        import base64

        sistem = self.giris("sistem")
        ogretmen = self.giris("deniz")
        once = {p.name for p in self.yedekler.glob("*.db")}
        yeni_dosyalar: set[str] = set()
        try:
            with tempfile.TemporaryDirectory() as dizin:
                kaynak = Path(dizin) / "yerel_dose.db"
                self.db.yedekle(kaynak)
                govde = {"dosya_adi": kaynak.name,
                         "icerik": base64.b64encode(kaynak.read_bytes()).decode("ascii")}
                yukle = self.c.post("/api/yedekler/yukle", headers=sistem, json=govde)
                self.assertEqual(yukle.status_code, 200, yukle.text)
                ad = yukle.json()["dosya_adi"]
                yeni_dosyalar.add(ad)
                self.assertTrue(ad.startswith("dose_yedek_yuklenen_"))
                self.assertIn(ad, {x["dosya_adi"] for x in self.c.get(
                    "/api/yedekler", headers=sistem).json()["yedekler"]})

                self.assertEqual(self.c.get("/api/yedekler", headers=ogretmen).status_code, 403)
                self.assertEqual(self.c.post("/api/yedekler", headers=ogretmen).status_code, 403)
                self.assertEqual(self.c.post("/api/yedekler/yukle", headers=ogretmen, json=govde).status_code, 403)
                gecersiz = self.c.post("/api/yedekler/yukle", headers=sistem,
                                       json={"dosya_adi": "bozuk.db",
                                             "icerik": base64.b64encode(b"x" * 120).decode("ascii")})
                self.assertEqual(gecersiz.status_code, 400)
                uyumsuz = self.c.post("/api/yedekler/yukle", headers=sistem,
                                      json={"dosya_adi": "metin.txt", "icerik": govde["icerik"]})
                self.assertEqual(uyumsuz.status_code, 400)
        finally:
            yeni_dosyalar |= {p.name for p in self.yedekler.glob("*.db")} - once
            for ad in yeni_dosyalar:
                (self.yedekler / ad).unlink(missing_ok=True)

    def test_denetim_formu_rol_paylasimi_cakisma_rapor_ve_senkron(self):
        mudur, tmy, ogretmen, sayman = (self.giris("ali"), self.giris("veli"), self.giris("deniz"),
                                       self.giris("ayse"))
        olustur = self.c.post("/api/denetimler", headers=mudur,
                              json={"baslik": "2026 Denetimi", "bilgiler": {"okul_adi": "Derik MTAL"}})
        self.assertEqual(olustur.status_code, 200, olustur.text)
        denetim_id = olustur.json()["id"]
        detay = self.c.get(f"/api/denetimler/{denetim_id}", headers=ogretmen).json()
        self.assertEqual(len(detay["maddeler"]), 17)
        self.assertEqual(detay["maddeler"][0]["durum"], "Değerlendirilmedi")

        ogretmen_madde = self.c.put(f"/api/denetimler/{denetim_id}/maddeler/3.1.01", headers=ogretmen,
                                    json={"surum": 1, "durum": "Uygun", "aciklama": "Uygulama kayıtları incelendi."})
        self.assertEqual(ogretmen_madde.status_code, 200, ogretmen_madde.text)
        self.assertEqual(ogretmen_madde.json()["surum"], 2)
        eksik_aciklama = self.c.put(f"/api/denetimler/{denetim_id}/maddeler/3.1.02", headers=ogretmen,
                                    json={"surum": 1, "durum": "Eksik"})
        self.assertEqual(eksik_aciklama.status_code, 400)
        self.assertEqual(self.c.put(f"/api/denetimler/{denetim_id}/maddeler/5.1.01", headers=ogretmen,
                                    json={"surum": 1, "durum": "Eksik"}).status_code, 403)
        self.assertEqual(self.c.post("/api/denetimler", headers=ogretmen,
                                     json={"baslik": "Yetkisiz"}).status_code, 403)

        mali_madde = self.c.put(f"/api/denetimler/{denetim_id}/maddeler/5.1.01", headers=sayman,
                                json={"surum": 1, "durum": "Eksik", "aciklama": "Mutabakat belgesi eksik."})
        self.assertEqual(mali_madde.status_code, 200, mali_madde.text)
        self.assertEqual(self.c.put(f"/api/denetimler/{denetim_id}/maddeler/3.1.01", headers=tmy,
                                    json={"surum": 1, "durum": "Eksik", "aciklama": "Eksik kayıt."}).status_code, 409)

        onceki_islem = self.db.son_islem_no()
        rapor_bilgisi = {"surum": 1, "baslik": "2026 Denetimi", "durum": "Taslak",
                         "bilgiler": {"okul_adi": "Derik MTAL", "giris": "Genel inceleme."}}
        self.assertEqual(self.c.put(f"/api/denetimler/{denetim_id}", headers=mudur, json=rapor_bilgisi).status_code, 200)
        cakisan = self.c.put(f"/api/denetimler/{denetim_id}", headers=tmy, json=rapor_bilgisi)
        self.assertEqual(cakisan.status_code, 409)

        ogrenci_senkron = self.c.get(f"/api/senkron?son_no={onceki_islem}", headers=ogretmen).json()
        self.assertNotIn("rol", ogrenci_senkron["degisiklikler"][0])
        degisiklik = self.c.get(f"/api/senkron?son_no={onceki_islem}", headers=sayman).json()
        self.assertTrue(any(x["tablo"] in ("denetim", "denetim_madde") for x in degisiklik["degisiklikler"]))
        self.assertEqual(self.c.get(f"/api/denetimler/{denetim_id}/rapor", headers=ogretmen).status_code, 403)
        rapor = self.c.get(f"/api/denetimler/{denetim_id}/rapor", headers=mudur)
        self.assertEqual(rapor.status_code, 200, rapor.text)
        import base64
        self.assertTrue(base64.b64decode(rapor.json()["icerik"]).startswith(b"%PDF"))

        tamamlanmamis = dict(rapor_bilgisi, surum=2, durum="Tamamlandı")
        sonuc = self.c.put(f"/api/denetimler/{denetim_id}", headers=mudur, json=tamamlanmamis)
        self.assertEqual(sonuc.status_code, 400)
        self.db.ex("UPDATE denetim_madde SET durum = 'Uygulanamaz' WHERE denetim_id = ?", (denetim_id,))
        eksik_rapor = self.c.put(f"/api/denetimler/{denetim_id}", headers=mudur, json=tamamlanmamis)
        self.assertEqual(eksik_rapor.status_code, 400, eksik_rapor.text)
        self.assertIn("Eksik alanlar", eksik_rapor.text)
        tamam_bilgileri = {
            alan: "0" if alan.endswith(("_kadin", "_erkek")) else "Yok"
            for alan in denetim_formu.BILGI_ALANLARI
        }
        tamam = self.c.put(f"/api/denetimler/{denetim_id}", headers=mudur,
                           json={**tamamlanmamis, "bilgiler": tamam_bilgileri})
        self.assertEqual(tamam.status_code, 200, tamam.text)
        self.assertEqual(self.c.get(f"/api/denetimler/{denetim_id}", headers=mudur).json()["durum"], "Tamamlandı")

    def test_yil_sonu_api_hesaplama_kayit_disari_aktarim_ve_yetki(self):
        import base64

        self.db.ex("DELETE FROM utp_kisi WHERE yil = ?", (YIL,))
        self.db.ex("DELETE FROM yil_sonu WHERE yil = ?", (YIL,))
        fis_id = self.db.ekle("yevmiye", {
            "fis_no": "YS-API-TEST", "tarih": f"{YIL}-12-31", "aciklama": "Yıl sonu API testi",
            "kaynak": "Manuel",
        })
        self.db.ekle("yevmiye_satir", {
            "yevmiye_id": fis_id, "hesap_kodu": "100", "aciklama": "Kasa", "borc": 0, "alacak": 10000,
        })
        self.db.ekle("yevmiye_satir", {
            "yevmiye_id": fis_id, "hesap_kodu": "600", "aciklama": "Gelir", "borc": 0, "alacak": 10000,
        })
        sayman = self.giris("ayse")
        ogretmen = self.giris("deniz")
        try:
            veri = self.c.get(f"/api/yil-sonu?yil={YIL}", headers=sayman)
            self.assertEqual(veri.status_code, 200, veri.text)
            self.assertTrue(veri.json()["yazabilir"])
            self.assertTrue(veri.json()["kisiler"])
            self.assertEqual(self.c.get(f"/api/yil-sonu?yil={YIL}", headers=ogretmen).status_code, 403)

            payload = {
                "donem_kari": max(10000, veri.json()["kar_onerilen"]),
                "faiz_duzeltme": 0,
                "utp_orani": 1 / 3,
                "kisiler": [{
                    "personel_id": kisi["id"], "unvan_grubu": kisi["unvan_grubu"],
                    "fiili_gun": 180, "fazla_calisma_saat": 0,
                    "katki_orani": 1, "takdir_orani": 0, "haric": False,
                    "haric_neden": "", "diger_isletme_utp": 0, "onceki_kumulatif": 0,
                } for kisi in veri.json()["kisiler"]],
            }
            hesap = self.c.post(f"/api/yil-sonu/{YIL}/hesapla", headers=sayman, json=payload)
            self.assertEqual(hesap.status_code, 200, hesap.text)
            self.assertEqual(hesap.json()["hatalar"], [])
            self.assertAlmostEqual(
                hesap.json()["ozet"]["dagitilan_utp"],
                sum(k["brut"] for k in hesap.json()["kisiler"]), places=2)

            kayit = self.c.post(f"/api/yil-sonu/{YIL}", headers=sayman, json=payload)
            self.assertEqual(kayit.status_code, 200, kayit.text)
            self.assertEqual(self.db.deger("SELECT COUNT(*) FROM utp_kisi WHERE yil = ?", (YIL,), 0),
                             len(payload["kisiler"]))
            for tur, baslangic in (
                    ("cetvel", b"%PDF"), ("banka", b"PK"),
                    ("merkez-yazisi", b"%PDF"), ("shcek-yazisi", b"%PDF")):
                disari = self.c.get(f"/api/yil-sonu/{YIL}/disa-aktar?tur={tur}", headers=sayman)
                self.assertEqual(disari.status_code, 200, disari.text)
                self.assertTrue(base64.b64decode(disari.json()["icerik"]).startswith(baslangic))

            hatali = {**payload, "utp_orani": 0.5}
            red = self.c.post(f"/api/yil-sonu/{YIL}", headers=sayman, json=hatali)
            self.assertEqual(red.status_code, 400, red.text)
            self.assertEqual(self.c.post(f"/api/yil-sonu/{YIL}/hesapla", headers=ogretmen,
                                         json=payload).status_code, 403)
            self.assertEqual(self.c.post(f"/api/yil-sonu/{YIL}/geri-al", headers=sayman).status_code, 409)
        finally:
            with self.db.islem():
                self.db.ex("DELETE FROM utp_kisi WHERE yil = ?", (YIL,))
                self.db.ex("DELETE FROM yil_sonu WHERE yil = ?", (YIL,))
                self.db.ex("DELETE FROM resmi_yazi WHERE konu IN (?, ?)",
                           (f"{YIL} Yılı Kâr Payı Aktarımı", f"{YIL} Yılı SHÇEK Payı"))
                self.db.ex("DELETE FROM yevmiye_satir WHERE yevmiye_id = ?", (fis_id,))
                self.db.sil("yevmiye", fis_id)

    def test_resmi_belgeler_api_yazi_crud_ogrenci_muvafakati_ve_liste_pdf(self):
        import base64

        sayman = self.giris("ayse")
        ogretmen = self.giris("deniz")
        ogrenci_id = self.db.ekle("ogrenci", {
            "ad_soyad": "Belge API Öğrencisi", "okul_no": "BELGE-1", "sinif": "11-A",
            "veli_adi": "Veli Test", "veli_telefon": "05000000000", "aktif": 1,
        })
        yazi_id = None
        try:
            liste = self.c.get("/api/belgeler", headers=sayman)
            self.assertEqual(liste.status_code, 200, liste.text)
            self.assertIn("Serbest yazı", liste.json()["sablonlar"])
            self.assertTrue(any(o["id"] == ogrenci_id for o in liste.json()["ogrenciler"]))
            self.assertEqual(self.c.get("/api/belgeler", headers=ogretmen).status_code, 403)

            yazi = {
                "tarih": f"{YIL}-10-29", "sayi": "2026/123", "konu": "API belge testi",
                "muhatap": "İLGİLİ MAKAMA", "ilgi": "", "metin": "Test metni.",
                "imza1": "Sayman\nSayman", "imza2": "Müdür\nOkul Müdürü",
            }
            ekle = self.c.post("/api/belgeler/yazilar", headers=sayman, json=yazi)
            self.assertEqual(ekle.status_code, 200, ekle.text)
            yazi_id = ekle.json()["id"]
            detay = self.c.get(f"/api/belgeler/yazilar/{yazi_id}", headers=sayman)
            self.assertEqual(detay.status_code, 200, detay.text)
            guncel = {**yazi, "konu": "Güncel belge testi"}
            self.assertEqual(self.c.put(f"/api/belgeler/yazilar/{yazi_id}", headers=sayman,
                                        json=guncel).status_code, 200)
            pdf = self.c.get(f"/api/belgeler/yazilar/{yazi_id}/pdf", headers=sayman)
            self.assertEqual(pdf.status_code, 200, pdf.text)
            self.assertTrue(base64.b64decode(pdf.json()["icerik"]).startswith(b"%PDF"))
            self.assertEqual(self.c.post("/api/belgeler/yazilar", headers=sayman,
                                         json={**yazi, "tarih": "not-a-date"}).status_code, 400)

            muvafakat = self.c.post("/api/belgeler/muvafakatname", headers=sayman,
                                    json={"ogrenci_idler": [ogrenci_id]})
            self.assertEqual(muvafakat.status_code, 200, muvafakat.text)
            self.assertTrue(base64.b64decode(muvafakat.json()["icerik"]).startswith(b"%PDF"))
            tekrar = self.c.post("/api/belgeler/muvafakatname", headers=sayman,
                                 json={"ogrenci_idler": [ogrenci_id, ogrenci_id]})
            self.assertEqual(tekrar.status_code, 400)
            for url in ("/api/belgeler/ogrenci-listesi", "/api/belgeler/personel-listesi"):
                disari = self.c.get(url, headers=sayman)
                self.assertEqual(disari.status_code, 200, disari.text)
                self.assertTrue(base64.b64decode(disari.json()["icerik"]).startswith(b"%PDF"))
        finally:
            with self.db.islem():
                if yazi_id:
                    self.db.sil("resmi_yazi", yazi_id)
                self.db.sil("ogrenci", ogrenci_id)

    def test_kullanici_yonetimi_yetki_hesap_ve_gecici_sifre_kontrollerini_uygular(self):
        admin = self.giris("sistem")
        mudur = self.giris("ali")
        self.assertEqual(self.c.get("/api/kullanicilar", headers=mudur).status_code, 403)
        liste = self.c.get("/api/kullanicilar", headers=admin)
        self.assertEqual(liste.status_code, 200, liste.text)
        satir = next(x for x in liste.json()["yetki_tablosu"] if x["sayfa"] == "Kullanıcılar")
        self.assertEqual(satir["roller"][M.ROL_SISTEM], "y")
        self.assertEqual(satir["roller"][M.ROL_MUDUR], "")

        payload = {"kullanici_adi": "flet.yonetim", "ad_soyad": "Flet Kullanıcı",
                   "rol": M.ROL_MEMUR, "personel_id": None, "bolum_id": None, "sifre": "Gecici123"}
        ekle = self.c.post("/api/kullanicilar", headers=admin, json=payload)
        self.assertEqual(ekle.status_code, 200, ekle.text)
        kisi_id = ekle.json()["id"]
        try:
            tekrar = self.c.post("/api/kullanicilar", headers=admin,
                                 json={**payload, "kullanici_adi": "FLET.YONETIM"})
            self.assertEqual(tekrar.status_code, 409)
            gecersiz = self.c.post("/api/kullanicilar", headers=admin,
                                   json={**payload, "kullanici_adi": "gecersiz.rol", "rol": "Bilinmeyen"})
            self.assertEqual(gecersiz.status_code, 400)
            guncelle = self.c.put(f"/api/kullanicilar/{kisi_id}", headers=admin,
                                  json={"ad_soyad": "Güncel Kullanıcı", "rol": M.ROL_MEMUR,
                                        "personel_id": None, "bolum_id": None, "aktif": True})
            self.assertEqual(guncelle.status_code, 200, guncelle.text)
            sifre = self.c.post(f"/api/kullanicilar/{kisi_id}/sifre", headers=admin,
                                json={"sifre": "YeniGecici123"})
            self.assertEqual(sifre.status_code, 200, sifre.text)
            self.assertEqual(self.c.post("/api/giris", json={"kullanici_adi": "flet.yonetim",
                                                             "sifre": "YeniGecici123",
                                                             "kurum_kodu": "765432"}).status_code, 200)
            pasif = self.c.put(f"/api/kullanicilar/{kisi_id}", headers=admin,
                               json={"ad_soyad": "Güncel Kullanıcı", "rol": M.ROL_MEMUR,
                                     "personel_id": None, "bolum_id": None, "aktif": False})
            self.assertEqual(pasif.status_code, 200, pasif.text)
            self.assertEqual(self.c.post("/api/giris", json={"kullanici_adi": "flet.yonetim",
                                                             "sifre": "YeniGecici123",
                                                             "kurum_kodu": "765432"}).status_code, 401)
            self.assertEqual(self.c.delete(f"/api/kullanicilar/{self.sistem_id}",
                                           headers=admin).status_code, 409)
        finally:
            with self.db.islem():
                self.db.sil("kullanici", kisi_id)


if __name__ == "__main__":
    unittest.main()
