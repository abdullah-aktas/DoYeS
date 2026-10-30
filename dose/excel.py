"""Excel dışa ve içe aktarma."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

_INCE = Side(style="thin", color="B7C3D0")


def disa_aktar(yol: Path, sayfalar: list[tuple[str, list, list]], baslik: str | None = None) -> Path:
    """sayfalar: [(sayfa_adı, başlıklar, satırlar)]"""
    wb = Workbook()
    wb.remove(wb.active)
    for ad, basliklar, satirlar in sayfalar:
        ws = wb.create_sheet(title=_sayfa_adi(ad))
        r0 = 1
        if baslik:
            ws.cell(row=1, column=1, value=baslik).font = Font(bold=True, size=12, color="1F4E79")
            r0 = 3
        for j, b in enumerate(basliklar, 1):
            c = ws.cell(row=r0, column=j, value=b)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F4E79")
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = Border(left=_INCE, right=_INCE, top=_INCE, bottom=_INCE)
        for i, s in enumerate(satirlar, r0 + 1):
            for j, v in enumerate(s, 1):
                c = ws.cell(row=i, column=j, value=v)
                c.border = Border(left=_INCE, right=_INCE, top=_INCE, bottom=_INCE)
                if isinstance(v, float):
                    c.number_format = '#,##0.00'
        for j, b in enumerate(basliklar, 1):
            uzunluk = max([len(str(b))] + [len(str(s[j - 1])) for s in satirlar[:300] if j - 1 < len(s) and s[j - 1] is not None])
            ws.column_dimensions[get_column_letter(j)].width = min(max(10, uzunluk + 2), 60)
        ws.freeze_panes = ws.cell(row=r0 + 1, column=1)
    yol = Path(yol)
    wb.save(str(yol))
    return yol


def _sayfa_adi(ad: str) -> str:
    for ch in '[]:*?/\\':
        ad = ad.replace(ch, " ")
    return ad[:31] or "Sayfa"


def oku(yol: Path) -> dict[str, list[list]]:
    """Tüm sayfaları hücre değerleri (hesaplanmış) olarak döndürür."""
    wb = load_workbook(str(yol), data_only=True, read_only=True)
    sonuc = {}
    for ws in wb.worksheets:
        sonuc[ws.title] = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    return sonuc


def _norm(s) -> str:
    tablo = str.maketrans("ÇĞİIÖŞÜçğıöşü", "CGIIOSUcgiosu")
    return "".join(ch for ch in str(s or "").translate(tablo).lower() if ch.isalnum())


def _eslesir(hn: str, aday: str) -> bool:
    """Kısa adaylar ('no', 'tc' gibi) yalnız tam eşleşir; 'S.No' sütunu 'Okul no' sanılmasın."""
    an = _norm(aday)
    return hn == an or (len(an) >= 4 and an in hn)


TCKN_BASLIKLARI = ["kimlik", "tc", "tcno", "tckn", "tc no"]
AD_BASLIKLARI = ["adı soyadı", "adı ve soyadı", "ad soyad", "adsoyad", "isim"]
PERSONEL_ICE_AKTARMA = {
    "sicil_no": ["personel no", "sicil", "id"], "tckn": TCKN_BASLIKLARI,
    "ad_soyad": AD_BASLIKLARI + ["personel adı"],
    "kadro_unvani": ["unvan", "görevi", "branş"], "iban": ["iban"], "telefon": ["telefon"],
}
OGRENCI_ICE_AKTARMA = {
    "okul_no": ["okul no", "okul numarası", "öğrenci no", "numara", "no"], "tckn": TCKN_BASLIKLARI,
    "ad_soyad": AD_BASLIKLARI + ["öğrenci adı"], "sinif": ["sınıf", "şube"],
    "iban": ["iban"], "_gorev": ["görevi", "unvan"],
}


def ogrenci_satiri(gorev) -> bool:
    return _norm(str(gorev or "")).startswith(("ogrenc", "kursiyer", "stajyer"))


def _baslik_eslesmesi(satir, alan_esleri) -> dict[str, int]:
    eslesme = {}
    for j, h in enumerate(satir):
        hn = _norm(h) if isinstance(h, str) else ""
        if not hn:
            continue
        for alan, adaylar in alan_esleri.items():
            if alan not in eslesme and any(_eslesir(hn, a) for a in adaylar):
                eslesme[alan] = j
                break
    return eslesme


AD_ALANLARI = ("ad_soyad", "unvan", "ad")


def tablo_bul(yol: Path, alan_esleri: dict[str, list[str]], ad_alani: str | None = None,
              tarama: int = 200) -> tuple[list[dict], str | None]:
    """Başlık satırını otomatik bulur; alan_esleri: {alan: [olası başlık parçaları]}. Tüm sayfalarda en çok alanın
    eşleştiği (ad alanını da içeren) başlık seçilir. Aynı sayfada alt alta birden çok tablo olabileceği için okuma
    TOPLAM satırında, yeni bir başlık satırında veya art arda iki adsız satırda durur. (kayıtlar, sayfa adı)"""
    ad_alani = ad_alani or next((a for a in AD_ALANLARI if a in alan_esleri), "")
    sayfalar = oku(yol)
    en_iyi = (0, None, None, None)
    for ad, satirlar in sayfalar.items():
        for i, satir in enumerate(satirlar[:tarama]):
            eslesme = _baslik_eslesmesi(satir, alan_esleri)
            if ad_alani in alan_esleri and ad_alani not in eslesme:
                continue
            if len(eslesme) > en_iyi[0]:
                en_iyi = (len(eslesme), ad, i, eslesme)
    if not en_iyi[1]:
        return [], None
    _, ad, bas, eslesme = en_iyi
    sol, sag = max(0, min(eslesme.values()) - 2), max(eslesme.values()) + 1
    kayitlar, bos = [], 0
    for s in sayfalar[ad][bas + 1:]:
        k = {alan: (s[j] if j < len(s) else None) for alan, j in eslesme.items()}
        aralik = list(s[sol:sag])
        if any(_norm(v).startswith(("toplam", "geneltoplam")) for v in aralik if isinstance(v, str)):
            break
        if len(_baslik_eslesmesi(aralik, alan_esleri)) >= 2:
            break
        isim = k.get(ad_alani) if ad_alani in eslesme else next((v for v in k.values() if v not in (None, "")), None)
        if isim in (None, "") or (isinstance(isim, str) and not isim.strip()):
            bos += 1
            if kayitlar and bos >= 2:
                break
            continue
        bos = 0
        if all(not (isinstance(v, str) and v.startswith("=")) for v in k.values()):
            kayitlar.append(k)
    return kayitlar, ad


def eslestir_ve_oku(yol: Path, alan_esleri: dict[str, list[str]]) -> list[dict]:
    return tablo_bul(yol, alan_esleri)[0]
